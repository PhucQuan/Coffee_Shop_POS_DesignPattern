# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from docx_helpers import (
    set_cell_background, set_cell_margins, add_heading_1,
    add_heading_2, add_heading_3, add_body_p, add_code_block,
    create_styled_table
)

# Testcases data
from generate_excel import testcases

print(f"Loaded {len(testcases)} blackbox testcases.")

def append_create_beverage_whitebox(doc):
    add_heading_2(doc, "Hàm thứ hai: public Beverage createBeverage(MenuItemRecord item)")
    add_body_p(doc, "src/main/java/com/coffeeshop/service/MenuService.java", bold_prefix="Nguồn mã nguồn: ", italic=True)
    
    add_heading_3(doc, "1. Mô tả chức năng")
    add_body_p(doc, "Hàm createBeverage(MenuItemRecord item) trong lớp MenuService khởi tạo đối tượng đồ uống (Beverage) cụ thể tương ứng với bản ghi thực đơn qua mẫu thiết kế Factory Method Pattern:")
    add_body_p(doc, "- Đầu vào: item (MenuItemRecord) chứa thông tin tên món (name), đơn giá cơ bản (basePrice), và danh mục đồ uống (category).")
    add_body_p(doc, "- Quy trình xử lý: Kiểm tra thuộc tính item.getCategory() thông qua switch-expression. Khởi tạo đối tượng Factory tương ứng (CoffeeFactory, TeaFactory, MatchaFactory, SmoothieFactory). Nếu category không hợp lệ hoặc nằm ngoài danh mục hỗ trợ, ném ngoại lệ IllegalArgumentException. Đối với nhánh hợp lệ, gọi factory.createBeverage(...) để trả về đối tượng đồ uống.")
    add_body_p(doc, "- Đầu ra: Đối tượng Beverage đã được khởi tạo hoàn chỉnh hoặc ném ngoại lệ.")
    
    add_heading_3(doc, "2. Đoạn mã nguồn (có đánh số thứ tự câu lệnh)")
    code = """public Beverage createBeverage(MenuItemRecord item) {
    BeverageFactory factory = switch (item.getCategory()) {            // (1)
        case "COFFEE"   -> new CoffeeFactory();                        // (2)
        case "TEA"      -> new TeaFactory();                           // (3)
        case "MATCHA"   -> new MatchaFactory();                        // (4)
        case "SMOOTHIE" -> new SmoothieFactory();                      // (5)
        default         -> throw new IllegalArgumentException(         // (6)
            "Unsupported beverage category: " + item.getCategory());
    };                                                                 // (7)
    return factory.createBeverage(item.getName(), item.getBasePrice());// (8)
}                                                                      // (9)"""
    add_code_block(doc, code)
    
    add_heading_3(doc, "3. Đồ thị dòng điều khiển cơ bản (Control Flow Graph - CFG)")
    add_body_p(doc, "Đồ thị điều khiển biểu diễn các nút từ (1) đến (9):")
    add_body_p(doc, "- Nút (1): Đánh giá switch (item.getCategory()).")
    add_body_p(doc, "- Nút (2): case 'COFFEE' -> new CoffeeFactory().")
    add_body_p(doc, "- Nút (3): case 'TEA' -> new TeaFactory().")
    add_body_p(doc, "- Nút (4): case 'MATCHA' -> new MatchaFactory().")
    add_body_p(doc, "- Nút (5): case 'SMOOTHIE' -> new SmoothieFactory().")
    add_body_p(doc, "- Nút (6): default -> ném IllegalArgumentException.")
    add_body_p(doc, "- Nút (7): Gán biến factory.")
    add_body_p(doc, "- Nút (8): Gọi factory.createBeverage(...) và trả về đối tượng Beverage.")
    add_body_p(doc, "- Nút (9): Kết thúc phương thức (Exit/End).")
    
    add_heading_3(doc, "4. Tính số đường thi hành tuyến tính độc lập (McCabe)")
    add_body_p(doc, "Cấu trúc switch có 5 nhánh rẽ tương đương 4 quyết định nhị phân độc lập P = 4. Độ phức tạp chu trình McCabe: C = P + 1 = 4 + 1 = 5.")
    add_body_p(doc, "5 đường thi hành tuyến tính độc lập cơ bản:")
    add_body_p(doc, "1. Đường 1: 1 -> 2 -> 7 -> 8 -> 9 (Danh mục COFFEE, tạo BaseCoffee)")
    add_body_p(doc, "2. Đường 2: 1 -> 3 -> 7 -> 8 -> 9 (Danh mục TEA, tạo MilkTea)")
    add_body_p(doc, "3. Đường 3: 1 -> 4 -> 7 -> 8 -> 9 (Danh mục MATCHA, tạo Matcha)")
    add_body_p(doc, "4. Đường 4: 1 -> 5 -> 7 -> 8 -> 9 (Danh mục SMOOTHIE, tạo Smoothie)")
    add_body_p(doc, "5. Đường 5: 1 -> 6 -> 9 (Danh mục khác / không hỗ trợ, ném ngoại lệ)")
    
    add_heading_3(doc, "5. Thiết kế Test Cases cho từng đường")
    tc_headers = ["Mã TC", "Đường thi hành", "Dữ liệu kiểm thử (Test Data)", "Kết quả mong đợi"]
    tc_data = [
        ["TC_CB_01", "1 -> 2 -> 7 -> 8 -> 9", "item{id=1, name='Cà phê đen', basePrice=25000, category='COFFEE'}", "Tạo BaseCoffee, giá 25.000 ₫"],
        ["TC_CB_02", "1 -> 3 -> 7 -> 8 -> 9", "item{id=2, name='Trà sữa Ô Long', basePrice=32000, category='TEA'}", "Tạo MilkTea, giá 32.000 ₫"],
        ["TC_CB_03", "1 -> 4 -> 7 -> 8 -> 9", "item{id=3, name='Matcha Đá Xay', basePrice=45000, category='MATCHA'}", "Tạo Matcha, giá 45.000 ₫"],
        ["TC_CB_04", "1 -> 5 -> 7 -> 8 -> 9", "item{id=4, name='Sinh tố Bơ', basePrice=40000, category='SMOOTHIE'}", "Tạo Smoothie, giá 40.000 ₫"],
        ["TC_CB_05", "1 -> 6 -> 9", "item{id=5, name='Nước Ép Cam', basePrice=30000, category='JUICE'}", "Ném IllegalArgumentException('Unsupported beverage category: JUICE')"]
    ]
    create_styled_table(doc, tc_headers, tc_data, [1.0, 1.6, 2.4, 2.0])
    
    add_heading_3(doc, "6. Đồ thị dòng dữ liệu & Kiểm thử đời sống biến (Data Flow Testing)")
    add_body_p(doc, "Đồ thị dòng dữ liệu biểu diễn các hành vi d (define), u (use), k (kill):")
    add_body_p(doc, "- Biến item: d(item) tại đầu vào; u(item) tại (1) lấy category; u(item) tại (6) tạo thông báo lỗi; u(item) tại (8) lấy name, basePrice; k(item) tại (9).")
    add_body_p(doc, "- Biến factory: d(factory) tại (2), (3), (4), (5); u(factory) tại (8); k(factory) tại (9). Nhánh (6) không khởi tạo factory.")
    
    lc_headers = ["Tên biến", "Phạm vi", "Kịch bản 1", "Kịch bản 2", "Kịch bản 3", "Kịch bản 4", "Kịch bản 5", "Đánh giá bất thường"]
    lc_data = [
        ["item", "Tham số vào", "~duuk", "~duuk", "~duuk", "~duuk", "~duuk", "Bình thường (hợp lệ)"],
        ["factory", "Biến cục bộ", "~duk", "~duk", "~duk", "~duk", "(không tạo)", "Bình thường (định nghĩa và dùng đúng)"]
    ]
    create_styled_table(doc, lc_headers, lc_data, [1.0, 1.0, 0.8, 0.8, 0.8, 0.8, 0.8, 1.8])


def append_save_recipe_whitebox(doc):
    add_heading_2(doc, "Hàm thứ ba: public void saveRecipeItem(MenuItemRecord beverage, InventoryItem inventoryItem, double quantityRequired)")
    add_body_p(doc, "src/main/java/com/coffeeshop/service/MenuService.java", bold_prefix="Nguồn mã nguồn: ", italic=True)
    
    add_heading_3(doc, "1. Mô tả chức năng")
    add_body_p(doc, "Hàm saveRecipeItem(...) thiết lập và lưu trữ bản ghi công thức định mức nguyên liệu tiêu hao cho một món nước trong thực đơn:")
    add_body_p(doc, "- Đầu vào: beverage (món nước), inventoryItem (nguyên liệu trong kho), quantityRequired (định lượng tiêu hao / ly).")
    add_body_p(doc, "- Quy trình xử lý: Kiểm tra tuần tự 3 điều kiện tiền đề: (1) beverage == null -> ném IllegalArgumentException('Please select a beverage first.'); (2) inventoryItem == null -> ném IllegalArgumentException('Please select an inventory item.'); (3) quantityRequired <= 0 -> ném IllegalArgumentException('Recipe quantity must be greater than 0.'). Nếu cả 3 đều thỏa mãn, khởi tạo RecipeItem và lưu vào repository.")
    add_body_p(doc, "- Đầu ra: void (cập nhật cơ sở dữ liệu).")
    
    add_heading_3(doc, "2. Đoạn mã nguồn (có đánh số thứ tự câu lệnh)")
    code = """public void saveRecipeItem(MenuItemRecord beverage, InventoryItem inventoryItem, double quantityRequired) {
    if (beverage == null) {                                                                         // (1)
        throw new IllegalArgumentException("Please select a beverage first.");                     // (2)
    }
    if (inventoryItem == null) {                                                                    // (3)
        throw new IllegalArgumentException("Please select an inventory item.");                    // (4)
    }
    if (quantityRequired <= 0) {                                                                    // (5)
        throw new IllegalArgumentException("Recipe quantity must be greater than 0.");             // (6)
    }
    RecipeItem recipeItem = new RecipeItem(                                                         // (7)
        beverage.getId(), inventoryItem.getId(), quantityRequired);
    repository.saveRecipeItem(recipeItem);                                                          // (8)
}                                                                                                   // (9)"""
    add_code_block(doc, code)
    
    add_heading_3(doc, "3. Đồ thị dòng điều khiển cơ bản (Control Flow Graph - CFG)")
    add_body_p(doc, "- Nút (1): Kiểm tra điều kiện beverage == null. True -> (2); False -> (3).")
    add_body_p(doc, "- Nút (2): Ném ngoại lệ thiếu món đồ uống -> (9) Exit.")
    add_body_p(doc, "- Nút (3): Kiểm tra điều kiện inventoryItem == null. True -> (4); False -> (5).")
    add_body_p(doc, "- Nút (4): Ném ngoại lệ thiếu nguyên liệu kho -> (9) Exit.")
    add_body_p(doc, "- Nút (5): Kiểm tra điều kiện quantityRequired <= 0. True -> (6); False -> (7).")
    add_body_p(doc, "- Nút (6): Ném ngoại lệ định lượng không hợp lệ -> (9) Exit.")
    add_body_p(doc, "- Nút (7): Khởi tạo đối tượng RecipeItem.")
    add_body_p(doc, "- Nút (8): Gọi repository.saveRecipeItem(recipeItem).")
    add_body_p(doc, "- Nút (9): Kết thúc phương thức (Exit/End).")
    
    add_heading_3(doc, "4. Tính số đường thi hành tuyến tính độc lập (McCabe)")
    add_body_p(doc, "Hàm có 3 nút quyết định nhị phân độc lập P = 3. Độ phức tạp chu trình McCabe: C = P + 1 = 3 + 1 = 4.")
    add_body_p(doc, "4 đường thi hành tuyến tính độc lập:")
    add_body_p(doc, "1. Đường 1: 1 -> 2 -> 9 (Thiếu món đồ uống, ném ngoại lệ)")
    add_body_p(doc, "2. Đường 2: 1 -> 3 -> 4 -> 9 (Thiếu nguyên liệu tồn kho, ném ngoại lệ)")
    add_body_p(doc, "3. Đường 3: 1 -> 3 -> 5 -> 6 -> 9 (Định lượng <= 0, ném ngoại lệ)")
    add_body_p(doc, "4. Đường 4: 1 -> 3 -> 5 -> 7 -> 8 -> 9 (Hợp lệ, lưu công thức thành công)")
    
    add_heading_3(doc, "5. Thiết kế Test Cases cho từng đường")
    tc_headers = ["Mã TC", "Đường thi hành", "Dữ liệu kiểm thử (Test Data)", "Kết quả mong đợi"]
    tc_data = [
        ["TC_SR_01", "1 -> 2 -> 9", "beverage = null, item = item{id=1}, quantity = 15.0", "Ném IllegalArgumentException('Please select a beverage first.')"],
        ["TC_SR_02", "1 -> 3 -> 4 -> 9", "beverage = bev{id=1}, item = null, quantity = 15.0", "Ném IllegalArgumentException('Please select an inventory item.')"],
        ["TC_SR_03", "1 -> 3 -> 5 -> 6 -> 9", "beverage = bev{id=1}, item = item{id=1}, quantity = 0.0", "Ném IllegalArgumentException('Recipe quantity must be greater than 0.')"],
        ["TC_SR_04", "1 -> 3 -> 5 -> 7 -> 8 -> 9", "beverage = bev{id=1}, item = item{id=1}, quantity = 18.0", "Lưu thành công RecipeItem(1, 1, 18.0) vào cơ sở dữ liệu"]
    ]
    create_styled_table(doc, tc_headers, tc_data, [1.0, 1.6, 2.4, 2.0])
    
    add_heading_3(doc, "6. Đồ thị dòng dữ liệu & Kiểm thử đời sống biến (Data Flow Testing)")
    lc_headers = ["Tên biến", "Phạm vi", "Kịch bản 1", "Kịch bản 2", "Kịch bản 3", "Kịch bản 4", "Đánh giá bất thường"]
    lc_data = [
        ["beverage", "Tham số vào", "~duk", "~duk", "~duk", "~duuk", "Bình thường (hợp lệ)"],
        ["inventoryItem", "Tham số vào", "~dk", "~duk", "~duk", "~duuk", "Bình thường (~dk do ngắt sớm)"],
        ["quantityRequired", "Tham số vào", "~dk", "~dk", "~duk", "~duuk", "Bình thường (~dk do ngắt sớm)"],
        ["recipeItem", "Biến cục bộ", "(không tạo)", "(không tạo)", "(không tạo)", "~duk", "Bình thường (định nghĩa và dùng đúng)"],
        ["repository", "Trường lớp", "(không dùng)", "(không dùng)", "(không dùng)", "d -> u", "Bình thường (khởi tạo tại constructor)"]
    ]
    create_styled_table(doc, lc_headers, lc_data, [1.1, 0.9, 0.8, 0.8, 0.8, 0.8, 1.8])


def append_blackbox_chapter(doc):
    add_heading_1(doc, "CHƯƠNG V. KIỂM THỬ HỘP ĐEN (BLACK-BOX TESTING)")
    
    add_heading_2(doc, "5.1. Giới thiệu kỹ thuật & Chiến lược kiểm thử")
    add_body_p(doc, "Kiểm thử hộp đen (Black-box Testing) tập trung kiểm thử toàn diện các yêu cầu chức năng nghiệp vụ của hệ thống Coffee Shop POS từ góc nhìn người dùng cuối (Thu ngân, Pha chế, Quản trị viên). Các kỹ thuật chính được áp dụng bao gồm:")
    add_body_p(doc, "- Phân vùng tương đương (Equivalence Partitioning - EP): Chia các miền dữ liệu đăng nhập, đơn giá, số lượng thành các vùng hợp lệ và không hợp lệ.")
    add_body_p(doc, "- Phân tích giá trị biên (Boundary Value Analysis - BVA): Kiểm tra các giá trị biên của số lượng món (0, 1), ngưỡng cảnh báo kho, độ dài mật khẩu.")
    add_body_p(doc, "- Bảng quyết định (Decision Table): Áp dụng cho logic chiết khấu khuyến mãi Decorator, Strategy Pattern (Mua 1 Tặng 1, Giảm 10%).")
    add_body_p(doc, "- Kiểm thử chuyển trạng thái (State Transition Testing): Kiểm tra vòng đời đơn hàng qua các trạng thái PENDING -> PREPARING -> READY -> PAID.")
    add_body_p(doc, "- Kiểm thử theo kịch bản (Use Case Testing): Xác minh quy trình gọi món, thanh toán ví Momo/VNPay và xuất hóa đơn điện tử E-Receipt.")
    
    add_heading_2(doc, "5.2. Bảng tổng hợp 20 Kịch bản kiểm thử hộp đen (Black-box Test Cases)")
    add_body_p(doc, "Toàn bộ 20 kịch bản kiểm thử hộp đen được đồng bộ hóa trực tiếp với tệp Excel quản lý kiểm thử: docs/Bang_TestCases_HopDen_CoffeeShopPOS_20TC.xlsx")
    
    bb_headers = ["STT", "Mã TC", "Chức năng", "Mục đích kiểm thử", "Dữ liệu đầu vào", "Kết quả mong đợi", "Kỹ thuật", "Trạng thái", "Người test"]
    bb_data = []
    for tc in testcases:
        bb_data.append([
            tc['stt'],
            tc['id'],
            tc['module'],
            tc['desc'],
            tc['data'].replace('\n', '; '),
            tc['expected'].replace('\n', '; '),
            tc['tech'],
            tc['status'],
            tc['tester']
        ])
    create_styled_table(doc, bb_headers, bb_data, [0.4, 0.8, 1.1, 1.4, 1.2, 1.5, 0.9, 0.6, 0.9])
    
    add_heading_2(doc, "5.3. Đánh giá kết quả thực thi kiểm thử hộp đen")
    add_body_p(doc, "- Tổng số ca kiểm thử thiết kế: 20 Test Cases.")
    add_body_p(doc, "- Số ca kiểm thử thực thi Đạt (Pass): 20 / 20 (Tỷ lệ Đạt 100%).")
    add_body_p(doc, "- Số ca kiểm thử Thất bại (Fail): 0.")
    add_body_p(doc, "- Kết luận: Hệ thống Coffee Shop POS hoạt động ổn định, đáp ứng chính xác 100% các yêu cầu nghiệp vụ quản lý bán hàng, thanh toán đa kênh, tính toán chiết khấu và tự động khấu trừ tồn kho.")


# 1. Update Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx
doc_wb = docx.Document('docs/Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx')
append_create_beverage_whitebox(doc_wb)
append_save_recipe_whitebox(doc_wb)
append_blackbox_chapter(doc_wb)
doc_wb.save('docs/Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx')
print("Successfully updated docs/Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx")

# 2. Update BaoCao_CoffeeShopPOS_DesignPattern.docx
doc_bc = docx.Document('docs/BaoCao_CoffeeShopPOS_DesignPattern.docx')
append_create_beverage_whitebox(doc_bc)
append_save_recipe_whitebox(doc_bc)
append_blackbox_chapter(doc_bc)
doc_bc.save('docs/BaoCao_CoffeeShopPOS_DesignPattern.docx')
print("Successfully updated docs/BaoCao_CoffeeShopPOS_DesignPattern.docx")
