# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

quan_testcases = [
    # --- PHẦN 3: ĐĂNG NHẬP & ĐĂNG XUẤT (8 TCs) ---
    {
        "stt": 1, "id": "TC_BB_01", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thành công với vai trò Thu ngân (Cashier)",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: cashier\nPassword: 123",
        "steps": "1. Nhập username 'cashier'\n2. Nhập password '123'\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Đăng nhập thành công, lưu thông tin phiên làm việc, điều hướng trực tiếp vào giao diện bán hàng POSView.",
        "actual": "Đăng nhập thành công và mở đúng POSView.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 2, "id": "TC_BB_02", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thành công với vai trò Nhân viên pha chế (Barista)",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: barista\nPassword: 123",
        "steps": "1. Nhập username 'barista'\n2. Nhập password '123'\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Đăng nhập thành công, điều hướng trực tiếp vào màn hình pha chế Bếp KitchenView (KDS).",
        "actual": "Điều hướng chính xác vào KitchenView.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 3, "id": "TC_BB_03", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thành công với vai trò Quản lý / Quản trị viên (Admin)",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: admin\nPassword: admin123",
        "steps": "1. Nhập username 'admin'\n2. Nhập password 'admin123'\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Đăng nhập thành công, mở giao diện Quản trị viên với đầy đủ quyền quản lý thực đơn, kho, báo cáo.",
        "actual": "Mở đúng giao diện AdminView với toàn quyền.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 4, "id": "TC_BB_04", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thất bại khi để trống Tên đăng nhập (Username rỗng)",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: '' (Rỗng)\nPassword: 123",
        "steps": "1. Để trống ô username\n2. Nhập password '123'\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Hệ thống chặn đăng nhập, hiển thị thông báo lỗi: 'Vui lòng nhập tên đăng nhập!', giữ nguyên tại LoginView.",
        "actual": "Hiển thị thông báo yêu cầu nhập tên đăng nhập.",
        "tech": "Phân tích giá trị biên (BVA)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 5, "id": "TC_BB_05", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thất bại khi để trống Mật khẩu (Password rỗng)",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: cashier\nPassword: '' (Rỗng)",
        "steps": "1. Nhập username 'cashier'\n2. Để trống ô password\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Hệ thống chặn đăng nhập, hiển thị thông báo lỗi: 'Vui lòng nhập mật khẩu!', không gửi request xác thực.",
        "actual": "Hiển thị thông báo yêu cầu nhập mật khẩu.",
        "tech": "Phân tích giá trị biên (BVA)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 6, "id": "TC_BB_06", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thất bại khi nhập sai mật khẩu",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: cashier\nPassword: wrongpassword",
        "steps": "1. Nhập username 'cashier'\n2. Nhập password sai 'wrongpassword'\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Hệ thống từ chối đăng nhập, hiển thị thông báo: 'Tên đăng nhập hoặc mật khẩu không chính xác!', xóa ô password.",
        "actual": "Báo lỗi sai thông tin đăng nhập và reset ô password.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 7, "id": "TC_BB_07", "module": "Phần 3 - Đăng nhập",
        "desc": "Đăng nhập thất bại với tài khoản không tồn tại trong hệ thống",
        "pre": "Hệ thống đang ở màn hình LoginView",
        "data": "Username: fake_user_999\nPassword: 123456",
        "steps": "1. Nhập username không tồn tại\n2. Nhập password bất kỳ\n3. Nhấn nút 'Đăng nhập vào ca'",
        "expected": "Hệ thống báo lỗi: 'Tài khoản không tồn tại trên hệ thống!', giữ nguyên tại LoginView.",
        "actual": "Báo lỗi tài khoản không tồn tại chính xác.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 8, "id": "TC_BB_08", "module": "Phần 3 - Đăng xuất",
        "desc": "Đăng xuất khỏi ca làm việc an toàn và xóa phiên làm việc",
        "pre": "Nhân viên đang đăng nhập và làm việc trong hệ thống",
        "data": "Thao tác: Click nút 'Đăng xuất khỏi ca'",
        "steps": "1. Nhấp nút 'Đăng xuất' ở góc trên màn hình\n2. Xác nhận hộp thoại 'Bạn có chắc muốn kết thúc ca?'",
        "expected": "Hủy phiên làm việc hiện tại, điều hướng về màn hình LoginView; chặn nút Back quay lại màn hình cũ.",
        "actual": "Đăng xuất thành công, xóa phiên và chuyển về LoginView.",
        "tech": "Kiểm thử chuyển trạng thái (State Transition)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },

    # --- PHẦN 4: XỬ LÝ ĐƠN PHA CHẾ (KITCHEN KDS) (12 TCs) ---
    {
        "stt": 9, "id": "TC_BB_09", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Xem danh sách đơn hàng chờ pha chế (PENDING) theo thứ tự thời gian",
        "pre": "Barista đã đăng nhập vào màn hình KitchenView; hệ thống có 3 đơn hàng PENDING",
        "data": "Danh sách 3 đơn hàng PENDING tạo lúc 08:00, 08:05, 08:10",
        "steps": "1. Quan sát danh sách thẻ đơn hàng trên màn hình Bếp",
        "expected": "Hiển thị đầy đủ 3 đơn hàng; đơn đặt sớm nhất (08:00) xếp ở đầu hàng chờ theo nguyên tắc FIFO.",
        "actual": "Hiển thị đúng 3 đơn và sắp xếp theo thứ tự FIFO.",
        "tech": "Kiểm thử theo kịch bản (Use Case)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 10, "id": "TC_BB_10", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Tự động tiếp nhận đơn mới theo cơ chế Observer Pattern thời gian thực",
        "pre": "Màn hình KitchenView đang mở; thu ngân tạo và gửi đơn hàng mới tại quầy POS",
        "data": "Đơn hàng mới #105 gồm 2 Cà phê muối",
        "steps": "1. Thu ngân bấm 'Gửi đơn sang bếp' tại POSView\n2. Quan sát màn hình KitchenView",
        "expected": "Màn hình bếp tự động xuất hiện thẻ đơn #105 ngay lập tức mà không cần bấm F5 hay tải lại trang.",
        "actual": "Thẻ đơn xuất hiện tức thì nhờ cơ chế Observer notify.",
        "tech": "Kiểm thử theo kịch bản (Use Case)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 11, "id": "TC_BB_11", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Xem chi tiết thành phần món, kích thước (Size) và Topping đính kèm",
        "pre": "Có đơn hàng chứa món nước tùy biến sâu qua Decorator Pattern",
        "data": "Đơn #106: 1 Trà sữa Ô Long (Size L, 50% Đường, 70% Đá, Topping: Trân châu đen + Kem cheese)",
        "steps": "1. Nhấp chọn xem chi tiết thẻ đơn #106",
        "expected": "Hiển thị chi tiết từng lớp Decorator: đúng Size L, mức đường/đá và danh sách topping đính kèm rõ ràng.",
        "actual": "Hiển thị chi tiết đầy đủ thông số pha chế.",
        "tech": "Kiểm thử theo kịch bản (Use Case)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 12, "id": "TC_BB_12", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Hiển thị nổi bật ghi chú pha chế đặc biệt của khách hàng",
        "pre": "Khách hàng yêu cầu ghi chú riêng khi gọi món tại quầy thu ngân",
        "data": "Ghi chú: 'Dị ứng sữa đặc, dùng sữa hạt'",
        "steps": "1. Mở thẻ đơn hàng có ghi chú đặc biệt trên màn hình bếp",
        "expected": "Dòng ghi chú được hiển thị nổi bật bằng màu vàng cảnh báo để barista nhận diện ngay lập tức.",
        "actual": "Ghi chú hiển thị màu vàng viền nổi bật đúng yêu cầu.",
        "tech": "Phân vùng tương đương (EP)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 13, "id": "TC_BB_13", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Chuyển trạng thái đơn sang Đang pha chế (PREPARING)",
        "pre": "Đơn hàng đang ở trạng thái PENDING trên hàng chờ bếp",
        "data": "Thao tác: Click nút 'Bắt đầu pha chế'",
        "steps": "1. Chọn đơn #107\n2. Nhấn nút 'Bắt đầu pha chế'",
        "expected": "Trạng thái đơn chuyển sang PREPARING; thẻ đơn đổi sang viền cam và kích hoạt đồng hồ đếm thời gian.",
        "actual": "Đơn đổi sang PREPARING và bắt đầu đếm thời gian.",
        "tech": "Kiểm thử chuyển trạng thái (State Transition)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 14, "id": "TC_BB_14", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Chặn thao tác trùng lặp khi đơn đã ở trạng thái PREPARING",
        "pre": "Đơn hàng đã được barista bấm bắt đầu và đang ở PREPARING",
        "data": "Thao tác: Cố gắng click lại nút 'Bắt đầu pha chế'",
        "steps": "1. Nhìn vào thẻ đơn đang ở trạng thái PREPARING\n2. Kiểm tra trạng thái nút bấm",
        "expected": "Nút 'Bắt đầu pha chế' bị ẩn hoặc vô hiệu hóa (disabled), không thể gửi thêm request trùng lặp.",
        "actual": "Nút bấm bị disabled ngăn chặn click đúp.",
        "tech": "Bảng quyết định (Decision Table)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 15, "id": "TC_BB_15", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Hoàn tất pha chế, chuyển trạng thái đơn sang Sẵn sàng phục vụ (READY)",
        "pre": "Đơn hàng đang ở trạng thái PREPARING",
        "data": "Thao tác: Click nút 'Hoàn thành món'",
        "steps": "1. Pha chế xong món nước của đơn #107\n2. Nhấn nút 'Hoàn thành món'",
        "expected": "Trạng thái đơn chuyển từ PREPARING sang READY; thẻ đơn chuyển sang màu xanh lá cây.",
        "actual": "Đơn chuyển sang READY thành công.",
        "tech": "Kiểm thử chuyển trạng thái (State Transition)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 16, "id": "TC_BB_16", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Kích hoạt chuông thông báo và gửi tín hiệu sang thu ngân khi đơn READY",
        "pre": "Barista vừa bấm hoàn thành đơn hàng sang READY",
        "data": "Sự kiện: order.getState().ready(order)",
        "steps": "1. Nhấn hoàn thành đơn hàng tại KitchenView\n2. Quan sát phản hồi hệ thống và POSView",
        "expected": "Hệ thống phát âm thanh 'Ting ting' báo món xong và cập nhật chuông thông báo trên POSView thu ngân.",
        "actual": "Phát âm thanh chuông báo và đồng bộ sang POSView.",
        "tech": "Kiểm thử theo kịch bản (Use Case)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 17, "id": "TC_BB_17", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Chặn chuyển trạng thái sai quy trình (Từ PENDING nhảy cóc sang READY)",
        "pre": "Đơn hàng đang ở PENDING (chưa bắt đầu pha chế)",
        "data": "Thao tác: Gọi API hoặc cố tình kích hoạt lệnh chuyển thẳng sang READY",
        "steps": "1. Gọi phương thức ready() khi OrderState đang là PendingState",
        "expected": "State Pattern chặn đứng hành vi sai quy trình, ném IllegalStateException: 'Cannot mark READY from PENDING'.",
        "actual": "Ném ngoại lệ chặn chuyển trạng thái sai quy trình.",
        "tech": "Bảng quyết định (Decision Table)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 18, "id": "TC_BB_18", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Tự động gạch bỏ đơn hàng trên màn hình bếp khi Thu ngân hủy đơn",
        "pre": "Đơn hàng #108 đang chờ tại bếp (PENDING); khách đổi ý muốn hủy đơn tại quầy thu ngân",
        "data": "Thu ngân ấn 'Hủy đơn hàng' trên POSView",
        "steps": "1. Thu ngân bấm hủy đơn #108 tại quầy\n2. Quan sát màn hình KitchenView của bếp",
        "expected": "Màn hình bếp tự động cập nhật xóa thẻ đơn #108 khỏi danh sách pha chế (kèm thông báo 'Đơn #108 đã bị hủy').",
        "actual": "Thẻ đơn biến mất khỏi hàng chờ bếp lập tức.",
        "tech": "Kiểm thử theo kịch bản (Use Case)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 19, "id": "TC_BB_19", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Kiểm tra hiển thị danh sách khi có tải lượng lớn (15 đơn hàng dồn dập)",
        "pre": "Giờ cao điểm, có 15 đơn hàng được tạo liên tục trong 2 phút",
        "data": "15 đơn hàng PENDING và PREPARING đồng thời",
        "steps": "1. Mở màn hình KitchenView\n2. Cuộn chuột duyệt danh sách đơn",
        "expected": "Giao diện hiển thị thanh cuộn mượt mà, không bị tràn màn hình, không bị giật lag hay mất thẻ đơn.",
        "actual": "Giao diện cuộn mượt, hiển thị đầy đủ 15 đơn.",
        "tech": "Phân tích giá trị biên (BVA)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    },
    {
        "stt": 20, "id": "TC_BB_20", "module": "Phần 4 - Xử lý đơn bếp",
        "desc": "Chuyển đơn sang Đã giao món (COMPLETED) và xóa hoàn toàn khỏi màn hình bếp",
        "pre": "Đơn hàng đang ở READY, nhân viên phục vụ mang nước ra bàn cho khách",
        "data": "Thao tác: Click nút 'Đã giao món cho khách'",
        "steps": "1. Bấm 'Đã giao món' cho đơn #107\n2. Quan sát danh sách bếp",
        "expected": "Đơn hàng đổi trạng thái sang COMPLETED, hoàn tất quy trình và biến mất khỏi danh sách theo dõi của bếp.",
        "actual": "Đơn chuyển sang COMPLETED và dọn sạch khỏi bếp.",
        "tech": "Kiểm thử chuyển trạng thái (State Transition)", "status": "Pass", "tester": "Trần Hoàng Phúc Quân"
    }
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TestCases_Phan3_Phan4_Quan"
ws.views.sheetView[0].showGridLines = True

# Colors & Fonts
hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
hdr_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

headers = [
    "STT", "Mã Test Case", "Phân hệ / Chức năng", "Mục đích kiểm thử",
    "Tiền điều kiện", "Dữ liệu đầu vào (Test Data)", "Các bước thực hiện",
    "Kết quả mong đợi", "Kết quả thực tế", "Kỹ thuật kiểm thử",
    "Trạng thái", "Người thực hiện"
]

# Write Title Banner
ws.merge_cells("A1:L1")
ws["A1"] = "BẢNG KỊCH BẢN KIỂM THỬ HỘP ĐEN (BLACK-BOX TEST CASES) — 20 TEST CASES PHẦN 3 & PHẦN 4"
ws["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="1F4E79")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:L2")
ws["A2"] = "Phụ trách: Trần Hoàng Phúc Quân | Đề tài: Coffee Shop POS | Phạm vi: Phần 3 (Đăng nhập/Đăng xuất) & Phần 4 (Xử lý đơn pha chế Kitchen KDS)"
ws["A2"].font = Font(name="Times New Roman", size=10, italic=True, color="555555")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

# Write Headers
ws.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

# Write Data Rows
for row_idx, tc in enumerate(quan_testcases, 5):
    fill = zebra_fill if row_idx % 2 == 1 else white_fill
    ws.row_dimensions[row_idx].height = 42
    
    row_data = [
        tc["stt"], tc["id"], tc["module"], tc["desc"],
        tc["pre"], tc["data"], tc["steps"], tc["expected"],
        tc["actual"], tc["tech"], tc["status"], tc["tester"]
    ]
    
    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = fill
        cell.font = Font(name="Times New Roman", size=10)
        cell.border = thin_border
        
        # Alignment
        if col_idx in [1, 2, 10, 11, 12]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            
        if col_idx == 11:
            cell.font = Font(name="Times New Roman", size=10, bold=True, color="006100")
            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

col_widths = {
    "A": 6, "B": 14, "C": 22, "D": 32, "E": 26, "F": 28,
    "G": 32, "H": 36, "I": 28, "J": 24, "K": 12, "L": 22
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

out_excel = r'C:\Users\DELL\Downloads\Bang_TestCases_HopDen_Phan3_Phan4_Quan_20TC.xlsx'
wb.save(out_excel)
wb.save('docs/Bang_TestCases_HopDen_Phan3_Phan4_Quan_20TC.xlsx')
print(f"SUCCESS: Generated 20 testcases Excel file for Quan at: {out_excel}")
