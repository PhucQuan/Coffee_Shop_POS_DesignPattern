# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from generate_excel import testcases

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_p(doc, text="", bold_prefix="", font_size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(font_size)
        r_pre.font.bold = True
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic
    return p

def add_code_line(doc, line):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(line)
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def add_img(doc, img_path, width_in=2.5, caption=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(img_path, width=Inches(width_in))
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(6)
        rc = p_cap.add_run(caption)
        rc.font.name = 'Times New Roman'
        rc.font.size = Pt(10.5)
        rc.font.italic = True

def add_styled_table(doc, headers, data, widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    
    # Headers
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        c.text = h
        set_cell_bg(c, 'D9E1F2')
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = True
            
    # Rows
    for r_idx, row_values in enumerate(data):
        for c_idx, val in enumerate(row_values):
            c = tbl.cell(r_idx + 1, c_idx)
            c.text = str(val)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            if c_idx in [0, 1, 7, 8]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                if c_idx == 7 and 'Pass' in str(val):
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x27, 0x6A, 0x3C)
                elif c_idx == 1:
                    r.font.bold = True
                    
    if widths:
        for row in tbl.rows:
            for idx, w in enumerate(widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

print("Update helper defined.")
