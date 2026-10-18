# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_helpers import (
    set_cell_background, set_cell_margins, add_heading_1,
    add_heading_2, add_heading_3, add_body_p, add_code_block,
    create_styled_table
)
from generate_excel import testcases

def add_centered_picture(doc, img_path, width_inches=2.8, caption=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def build_full_report_file(target_file):
    doc = docx.Document('docs/BaoCao_CoffeeShopPOS_DesignPattern_backup.docx')
    
    # 1. Append White-box: createBeverage
    add_heading_2(doc, "Hàm thứ bảy: public Beverage createBeverage(MenuItemRecord item)")
    add_body_p(doc, "src/main/java/com/coffeeshop/service/MenuService.java", bold_prefix="Nguồn mã nguồn: ", italic=True)
    
    add_heading_3(doc, "1. Mô tả chức năng")
    add_body_p(doc, "Hàm createBeverage(MenuItemRecord item) trong lớp MenuService khởi tạo đối tượng đồ uống (Beverage) cụ thể tương ứng với bản ghi thực đơn qua mẫu thiết kế Factory Method Pattern:")
    add_body_p(doc, "- Đầu vào: item (MenuItemRecord) chứa thông tin tên món (name), đơn giá cơ bản (basePrice), và danh mục đồ uống (category).")
    add_body_p(doc, "- Quy trình xử lý: Kiểm tra thuộc tính item.getCategory() thông qua switch-expression. Khởi tạo đối tượng Factory tương ứng (CoffeeFactory, TeaFactory, MatchaFactory, SmoothieFactory). Nếu category không hợp lệ hoặc nằm ngoài danh mục hỗ trợ, ném ngoại lệ IllegalArgumentException. Đối với nhánh hợp lệ, gọi factory.createBeverage(...) để trả về đối tượng đồ uống.")
    add_body_p(doc, "- Đầu ra: Đối tượng Beverage đã được khởi tạo hoàn chỉnh hoặc ném ngoại lệ.")
    
    add_heading_3(doc, "2. Đoạn mã nguồn (có đánh số thứ tự câu lệnh)")
    code_bev = """public Beverage createBeverage(MenuItemRecord item) {
    BeverageFactory factory = switch (item.getCategory()) {            // (1)
        case "COFFEE"   -> new CoffeeFactory();                        // (2)
        case "TEA"      -> new TeaFactory();                           // (3)
        case "MATCHA"   -> new MatchaFactory();                        // (4)
        case "SMOOTHIE" -> new SmoothieFactory();                      // (5)
        default         -> throw new IllegalArgumentException(         // (6)
            "Unsupported beverage category: " + item.getCategory());
    };                                                                 // (7)
    return factory.createBeverage(item.getName(), item.getBasePrice());// (8)
}                                                                      // (9)"""
    add_code_block(doc, code_bev)
    
    add_heading_3(doc, "3. Lập đồ thị dòng điều khiển cơ bản (Control Flow Graph - CFG)")
    add_centered_picture(doc, 'docs/diagrams/cfg_create_beverage.png', width_inches=2.8, caption="Hình: Đồ thị dòng điều khiển cơ bản hàm createBeverage()")
    add_body_p(doc, "Đồ thị điều khiển biểu diễn các nút từ (1) đến (9):")
    add_body_p(doc, "- Nút (1): Đánh giá switch (item.getCategory()).")
    add_body_p(doc, "- Nút (2): case 'COFFEE' -> new CoffeeFactory().")
    add_body_p(doc, "- Nút (3): case 'TEA' -> new TeaFactory().")
    add_body_p(doc, "- Nút (4): case 'MATCHA' -> new MatchaFactory().")
    add_body_p(doc, "- Nút (5): case 'SMOOTHIE' -> new SmoothieFactory().")
    add_body_p(doc, "- Nút (6): default -> ném IllegalArgumentException.")
    add_body_p(doc, "- Nút (7): Nút gom gán biến factory.")
    add_body_p(doc, "- Nút (8): Gọi factory.createBeverage(...) và trả về đối tượng Beverage.")
    add_body_p(doc, "- Nút (9): Kết thúc phương thức (Exit/End).")
    
    add_heading_3(doc, "4. Tính số đường thi hành tuyến tính độc lập (McCabe)")
    add_body_p(doc, "Cấu trúc switch có 5 nhánh rẽ tương đương 4 quyết định nhị phân độc lập P = 4. Độ phức tạp chu trình McCabe: C = P + 1 = 4 + 1 = 5.")
    add_body_p(doc, "5 đường thi hành tuyến tính độc lập cơ bản:")
    add_body_p(doc, "1. 1 -> 2 -> 7 -> 8 -> 9 (Danh mục COFFEE)")
    add_body_p(doc, "2. 1 -> 3 -> 7 -> 8 -> 9 (Danh mục TEA)")
    add_body_p(doc, "3. 1 -> 4 -> 7 -> 8 -> 9 (Danh mục MATCHA)")
    add_body_p(doc, "4. 1 -> 5 -> 7 -> 8 -> 9 (Danh mục SMOOTHIE)")
    add_body_p(doc, "5. 1 -> 6 -> 9 (Danh mục không hợp lệ, ném ngoại lệ)")
    
    add_heading_3(doc, "5. Thiết kế Test Cases cho từng đường")
    tc_headers = ["STT", "Đường thi hành", "Dữ liệu kiểm thử (Test Data)", "Kết quả mong đợi"]
    tc_data = [
        ["1", "1 -> 2 -> 7 -> 8 -> 9", "item{id=1, name='Cà phê đen', basePrice=25000, category='COFFEE'}", "Tạo BaseCoffee, giá 25.000 ₫"],
        ["2", "1 -> 3 -> 7 -> 8 -> 9", "item{id=2, name='Trà sữa Ô Long', basePrice=32000, category='TEA'}", "Tạo MilkTea, giá 32.000 ₫"],
        ["3", "1 -> 4 -> 7 -> 8 -> 9", "item{id=3, name='Matcha Đá Xay', basePrice=45000, category='MATCHA'}", "Tạo Matcha, giá 45.000 ₫"],
        ["4", "1 -> 5 -> 7 -> 8 -> 9", "item{id=4, name='Sinh tố Bơ', basePrice=40000, category='SMOOTHIE'}", "Tạo Smoothie, giá 40.000 ₫"],
        ["5", "1 -> 6 -> 9", "item{id=5, name='Nước Ép Cam', basePrice=30000, category='JUICE'}", "Ném IllegalArgumentException('Unsupported beverage category: JUICE')"]
    ]
    create_styled_table(doc, tc_headers, tc_data, [0.6, 1.8, 2.4, 2.0])
    
    add_heading_3(doc, "6. Đồ thị dòng dữ liệu chung (DFG)")
    add_centered_picture(doc, 'docs/diagrams/dfg_create_beverage.png', width_inches=3.0, caption="Hình: Đồ thị dòng dữ liệu chung hàm createBeverage()")
    
    add_heading_3(doc, "7. Kiểm thử đời sống của từng biến")
    add_body_p(doc, "a. Kiểm thử đời sống biến item:")
    add_centered_picture(doc, 'docs/diagrams/var_item_create_beverage.png', width_inches=2.4, caption="Hình: Đời sống biến item")
    add_body_p(doc, "Kịch bản 1: ~duuk")
    add_body_p(doc, "2: ~duuk (giống 1)")
    add_body_p(doc, "3: ~duuk (giống 1)")
    add_body_p(doc, "4: ~duuk (giống 1)")
    add_body_p(doc, "5: ~duuk (dùng tại 1 và 6)")
    add_body_p(doc, "Cả 5 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")
    
    add_body_p(doc, "b. Kiểm thử đời sống biến factory:")
    add_centered_picture(doc, 'docs/diagrams/var_factory_create_beverage.png', width_inches=2.4, caption="Hình: Đời sống biến factory")
    add_body_p(doc, "Kịch bản 1: ~duk")
    add_body_p(doc, "2: ~duk (giống 1)")
    add_body_p(doc, "3: ~duk (giống 1)")
    add_body_p(doc, "4: ~duk (giống 1)")
    add_body_p(doc, "5: không được khởi tạo")
    add_body_p(doc, "Cả 5 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")
    
    # 2. Append White-box: saveRecipeItem
    add_heading_2(doc, "Hàm thứ tám: public void saveRecipeItem(MenuItemRecord beverage, InventoryItem inventoryItem, double quantityRequired)")
    add_body_p(doc, "src/main/java/com/coffeeshop/service/MenuService.java", bold_prefix="Nguồn mã nguồn: ", italic=True)
    
    add_heading_3(doc, "1. Mô tả chức năng")
    add_body_p(doc, "Hàm saveRecipeItem(...) thiết lập và lưu trữ bản ghi công thức định mức nguyên liệu tiêu hao cho một món nước trong thực đơn:")
    add_body_p(doc, "- Đầu vào: beverage (món nước), inventoryItem (nguyên liệu trong kho), quantityRequired (định lượng tiêu hao / 1 ly).")
    add_body_p(doc, "- Quy trình xử lý: Kiểm tra tuần tự 3 điều kiện tiền đề: (1) beverage == null; (2) inventoryItem == null; (3) quantityRequired <= 0. Nếu vi phạm bất kỳ điều kiện nào, ném ngoại lệ IllegalArgumentException tương ứng. Nếu tất cả hợp lệ, tạo RecipeItem và lưu vào cơ sở dữ liệu.")
    add_body_p(doc, "- Đầu ra: void (cập nhật cơ sở dữ liệu).")
    
    add_heading_3(doc, "2. Đoạn mã nguồn (có đánh số thứ tự câu lệnh)")
    code_rec = """public void saveRecipeItem(MenuItemRecord beverage, InventoryItem inventoryItem, double quantityRequired) {
    if (beverage == null) {                                                                         // (1)
        throw new IllegalArgumentException("Please select a beverage first.");                     // (2)
    }
    if (inventoryItem == null) {                                                                    // (3)
        throw new IllegalArgumentException("Please select an inventory item.");                    // (4)
    }
    if (quantityRequired <= 0) {                                                                    // (5)
        throw new IllegalArgumentException("Recipe quantity must be greater than 0.");             // (6)
    }
    RecipeItem recipeItem = new RecipeItem(                                                         // (7)
        beverage.getId(), inventoryItem.getId(), quantityRequired);
    repository.saveRecipeItem(recipeItem);                                                          // (8)
}                                                                                                   // (9)"""
    add_code_block(doc, code_rec)
    
    add_heading_3(doc, "3. Lập đồ thị dòng điều khiển cơ bản (Control Flow Graph - CFG)")
    add_centered_picture(doc, 'docs/diagrams/cfg_save_recipe.png', width_inches=2.4, caption="Hình: Đồ thị dòng điều khiển cơ bản hàm saveRecipeItem()")
    add_body_p(doc, "- Nút (1): Kiểm tra beverage == null. True -> (2); False -> (3).")
    add_body_p(doc, "- Nút (2): Ném ngoại lệ thiếu món đồ uống -> (9) Exit.")
    add_body_p(doc, "- Nút (3): Kiểm tra inventoryItem == null. True -> (4); False -> (5).")
    add_body_p(doc, "- Nút (4): Ném ngoại lệ thiếu nguyên liệu kho -> (9) Exit.")
    add_body_p(doc, "- Nút (5): Kiểm tra quantityRequired <= 0. True -> (6); False -> (7).")
    add_body_p(doc, "- Nút (6): Ném ngoại lệ định lượng không hợp lệ -> (9) Exit.")
    add_body_p(doc, "- Nút (7): Khởi tạo đối tượng RecipeItem.")
    add_body_p(doc, "- Nút (8): Gọi repository.saveRecipeItem(recipeItem).")
    add_body_p(doc, "- Nút (9): Kết thúc phương thức (Exit/End).")
    
    add_heading_3(doc, "4. Tính số đường thi hành tuyến tính độc lập (McCabe)")
    add_body_p(doc, "Hàm có 3 nút quyết định nhị phân độc lập P = 3. Độ phức tạp chu trình McCabe: C = P + 1 = 3 + 1 = 4.")
    add_body_p(doc, "4 đường thi hành tuyến tính độc lập:")
    add_body_p(doc, "1. 1 -> 2 -> 9 (Thiếu món đồ uống)")
    add_body_p(doc, "2. 1 -> 3 -> 4 -> 9 (Thiếu nguyên liệu tồn kho)")
    add_body_p(doc, "3. 1 -> 3 -> 5 -> 6 -> 9 (Định lượng tiêu hao <= 0)")
    add_body_p(doc, "4. 1 -> 3 -> 5 -> 7 -> 8 -> 9 (Hợp lệ, lưu công thức thành công)")
    
    add_heading_3(doc, "5. Thiết kế Test Cases cho từng đường")
    tc_data_rec = [
        ["1", "1 -> 2 -> 9", "beverage = null, item = item{id=1}, quantity = 15.0", "Ném IllegalArgumentException('Please select a beverage first.')"],
        ["2", "1 -> 3 -> 4 -> 9", "beverage = bev{id=1}, item = null, quantity = 15.0", "Ném IllegalArgumentException('Please select an inventory item.')"],
        ["3", "1 -> 3 -> 5 -> 6 -> 9", "beverage = bev{id=1}, item = item{id=1}, quantity = 0.0", "Ném IllegalArgumentException('Recipe quantity must be greater than 0.')"],
        ["4", "1 -> 3 -> 5 -> 7 -> 8 -> 9", "beverage = bev{id=1}, item = item{id=1}, quantity = 18.0", "Lưu thành công RecipeItem(1, 1, 18.0) vào cơ sở dữ liệu"]
    ]
    create_styled_table(doc, tc_headers, tc_data_rec, [0.6, 1.8, 2.4, 2.0])
    
    add_heading_3(doc, "6. Đồ thị dòng dữ liệu chung (DFG)")
    add_centered_picture(doc, 'docs/diagrams/dfg_save_recipe.png', width_inches=3.0, caption="Hình: Đồ thị dòng dữ liệu chung hàm saveRecipeItem()")
    
    add_heading_3(doc, "7. Kiểm thử đời sống của từng biến")
    add_body_p(doc, "a. Kiểm thử đời sống biến beverage:")
    add_centered_picture(doc, 'docs/diagrams/var_beverage_save_recipe.png', width_inches=2.2, caption="Hình: Đời sống biến beverage")
    add_body_p(doc, "Kịch bản 1: ~duk")
    add_body_p(doc, "2: ~duk (giống 1)")
    add_body_p(doc, "3: ~duk (giống 1)")
    add_body_p(doc, "4: ~duuk")
    add_body_p(doc, "Cả 4 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")
    
    add_body_p(doc, "b. Kiểm thử đời sống biến inventoryItem:")
    add_centered_picture(doc, 'docs/diagrams/var_inventory_save_recipe.png', width_inches=2.2, caption="Hình: Đời sống biến inventoryItem")
    add_body_p(doc, "Kịch bản 1: ~dk (ngắt sớm do ngoại lệ tiền đề)")
    add_body_p(doc, "2: ~duk")
    add_body_p(doc, "3: ~duk (giống 2)")
    add_body_p(doc, "4: ~duuk")
    add_body_p(doc, "Cả 4 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")

    # 3. Append Chapter 5: BLACKBOX
    add_heading_1(doc, "CHƯƠNG 5: KIỂM THỬ HỘP ĐEN")
    
    add_heading_2(doc, "5.1. Tổng hợp các Testcases")
    add_body_p(doc, "Kiểm thử hộp đen (Black-box Testing) được tiến hành nhằm xác minh tính đúng đắn của toàn bộ các yêu cầu chức năng nghiệp vụ của hệ thống Coffee Shop POS từ góc nhìn người dùng cuối (Thu ngân, Nhân viên pha chế, Quản trị viên).")
    add_body_p(doc, "Bộ kịch bản kiểm thử gồm 20 Test Cases phủ kín 20 chức năng cốt lõi (từ F01 đến F20), kết hợp đồng thời các kỹ thuật chuẩn mực:")
    add_body_p(doc, "- Phân vùng tương đương (Equivalence Partitioning - EP): Phân chia miền giá trị hợp lệ/không hợp lệ cho đăng nhập, giá tiền, số lượng.")
    add_body_p(doc, "- Phân tích giá trị biên (Boundary Value Analysis - BVA): Đánh giá hành vi tại các điểm ranh giới (số lượng = 0, tồn kho chạm mức cảnh báo reorderLevel, độ dài mật khẩu tối thiểu).")
    add_body_p(doc, "- Bảng quyết định (Decision Table): Kiểm tra các tổ hợp điều kiện phức tạp của Decorator Pattern (tùy biến size, topping) và Strategy Pattern (chiết khấu phần trăm, Mua 1 Tặng 1).")
    add_body_p(doc, "- Kiểm thử chuyển trạng thái (State Transition Testing): Xác thực vòng đời đơn hàng theo mẫu State Pattern (PENDING -> PREPARING -> READY -> PAID) và chặn hủy đơn không hợp lệ.")
    add_body_p(doc, "- Kiểm thử theo kịch bản (Use Case Testing): Kiểm thử quy trình bán hàng khép kín (End-to-End) từ tạo đơn, thanh toán ví Momo/VNPay đến kết xuất hóa đơn điện tử E-Receipt.")

    add_body_p(doc, "Bảng 5.1: Tổng hợp 20 kịch bản kiểm thử hộp đen chuẩn hóa (đồng bộ với file Excel: docs/Bang_TestCases_HopDen_CoffeeShopPOS_20TC.xlsx):", italic=True)

    bb_headers = ["STT", "Mã TC", "Chức năng", "Mục đích kiểm thử", "Dữ liệu đầu vào", "Kết quả mong đợi", "Kỹ thuật", "Trạng thái", "Người test"]
    bb_data = []
    for tc in testcases:
        bb_data.append([
            tc['stt'],
            tc['id'],
            tc['module'],
            tc['desc'],
            tc['data'].replace('\n', '; '),
            tc['expected'].replace('\n', '; '),
            tc['tech'],
            tc['status'],
            tc['tester']
        ])
    create_styled_table(doc, bb_headers, bb_data, [0.4, 0.8, 1.1, 1.4, 1.2, 1.5, 0.9, 0.6, 0.9])

    add_heading_2(doc, "5.2. Quản lý bug và báo cáo lỗi")
    add_body_p(doc, "Trong quá trình thực thi kiểm thử hộp đen, nhóm vận dụng quy trình quản lý vòng đời lỗi phần mềm nhằm phát hiện, phân loại và khắc phục kịp thời các sai sót trước khi nghiệm thu.")
    add_body_p(doc, "Phân loại mức độ nghiêm trọng của lỗi (Severity):")
    add_body_p(doc, "- Blocker / Critical: Lỗi khiến hệ thống ngừng hoạt động, sai lệch dữ liệu tài chính hoặc thanh toán thất bại nhưng trạng thái vẫn chuyển thành công.")
    add_body_p(doc, "- Major: Lỗi logic nghiệp vụ không đúng theo đặc tả (ví dụ: công thức giảm giá tính sai phụ phí topping, đơn hàng PAID vẫn cho phép ấn nút Hủy).")
    add_body_p(doc, "- Minor / Trivial: Lỗi giao diện hiển thị, căn lề lệch trên modal hoặc định dạng số tiền chưa có dấu phân cách hàng nghìn (₫).")

    add_body_p(doc, "Bảng 5.2: Nhật ký các lỗi phát hiện trong quá trình kiểm thử hộp đen:", italic=True)
    bug_headers = ["Mã Bug", "Chức năng", "Mô tả lỗi phát hiện", "Mức độ", "Nguyên nhân gốc", "Tình trạng"]
    bug_data = [
        ["BUG-01", "F05 - Tính tiền", "Khi áp dụng giảm giá Mua 1 Tặng 1, tổng tiền bị âm nếu chọn món tặng có giá cao hơn món mua.", "Major", "Chưa dùng hàm Math.max(0, subtotal - discount) trong hàm recalculate.", "Đã khắc phục (Closed)"],
        ["BUG-02", "F07 - Trạng thái", "Nhân viên vẫn có thể ấn hủy đơn hàng khi đơn đã ở trạng thái PAID.", "Critical", "Trạng thái PaidState chưa override phương thức cancel() để ném InvalidStateTransitionException.", "Đã khắc phục (Closed)"],
        ["BUG-03", "F17 - Quản lý kho", "Kho nguyên liệu bị trừ 2 lần khi thu ngân ấn nút thanh toán nhiều lần liên tiếp.", "Critical", "Thiếu cờ kiểm tra order.isInventoryDeducted() trước khi thực hiện deductForOrder.", "Đã khắc phục (Closed)"],
        ["BUG-04", "F16 - Tùy biến món", "Modal tùy biến món không hiển thị topping đã bị quản trị viên tắt (active=false).", "Minor", "Query lấy danh sách topping chưa lọc điều kiện active == true.", "Đã khắc phục (Closed)"],
        ["BUG-05", "F08 - Thanh toán", "Mã giao dịch ví điện tử Momo bị trùng khi phát sinh nhiều đơn hàng cùng lúc.", "Major", "Hàm giả lập mã giao dịch dùng chuỗi ngẫu nhiên quá ngắn.", "Đã khắc phục (Closed)"]
    ]
    create_styled_table(doc, bug_headers, bug_data, [0.8, 1.2, 2.4, 0.8, 1.8, 1.0])
    add_body_p(doc, "Đánh giá chất lượng: 100% các lỗi ghi nhận đều đã được đội ngũ phát triển khắc phục triệt để và kiểm thử hồi quy (Regression Testing) thành công, đảm bảo hệ thống đạt độ ổn định tuyệt đối.")

    add_heading_2(doc, "5.3. Triển khai Kiểm thử Tự động")
    add_body_p(doc, "Để nâng cao hiệu quả kiểm thử và phục vụ quá trình kiểm thử hồi quy liên tục, nhóm nghiên cứu đã xây dựng bộ kịch bản kiểm thử tự động toàn diện sử dụng framework Playwright (Node.js/JavaScript).")
    add_body_p(doc, "Kiến trúc và tính năng của bộ kiểm thử Playwright:")
    add_body_p(doc, "- Tự động hóa điều hướng và tương tác với giao diện người dùng web POS (http://localhost:5173 và http://localhost:5174).")
    add_body_p(doc, "- Chạy kiểm thử tự động 1-click thông qua script run-e2e-tests.bat.")
    add_body_p(doc, "- Tự động quay video màn hình HD (1280x720) cho từng kịch bản kiểm thử và lưu trữ tại thư mục e2e-tests/videos/.")
    add_body_p(doc, "- Tự động xuất báo cáo kiểm thử HTML trực quan (Playwright HTML Report) tích hợp ảnh chụp màn hình và video xem lại thao tác.")

    add_body_p(doc, "Bảng 5.3: Danh mục kịch bản kiểm thử tự động Playwright & Video ghi hình:", italic=True)
    auto_headers = ["Mã Kịch Bản", "Nội dung kiểm thử tự động", "Design Pattern liên quan", "Video ghi hình", "Kết quả"]
    auto_data = [
        ["TC01", "Kiểm tra giao diện thu ngân, tìm kiếm món thời gian thực & lọc theo từng danh mục", "UI / Filter Strategy", "TC01_GiaoDien_TimKiem_DanhMuc.webm", "Pass (100%)"],
        ["TC02", "Tùy biến món nước chuyên sâu: Size L (+10.000₫), Topping, mức đường, đá & thêm giỏ hàng", "Decorator Pattern", "TC02_TuyBienMon_Decorator_ThemVaoGio.webm", "Pass (100%)"],
        ["TC03", "Áp dụng các chiến lược khuyến mãi: Giảm giá 10%, VIP Fixed 15.000₫, Mua 1 Tặng 1 (BOGO)", "Strategy Pattern", "TC03_KhuyenMai_StrategyPattern.webm", "Pass (100%)"],
        ["TC04", "Quy trình thanh toán đa kênh: Ví điện tử Momo, VNPay QR Code và xuất hóa đơn điện tử E-Receipt", "Adapter Pattern", "TC04_ThanhToan_AdapterPattern_HoaDon.webm", "Pass (100%)"],
        ["TC05", "Đồng bộ màn hình pha chế (Kitchen KDS), luân chuyển trạng thái và tổng hợp báo cáo doanh thu", "State & Observer Pattern", "TC05_KitchenKDS_StatePattern_Admin.webm", "Pass (100%)"]
    ]
    create_styled_table(doc, auto_headers, auto_data, [0.8, 2.5, 1.4, 2.2, 0.9])

    # 4. Append Chapter 6: JIRA
    add_heading_1(doc, "CHƯƠNG 6: HỆ THỐNG QUẢN LÝ BUG")
    add_body_p(doc, "Hệ thống quản lý lỗi trong dự án được tổ chức bài bản trên công cụ Jira kết hợp biểu mẫu báo cáo chuẩn hóa. Mỗi lỗi được theo dõi chặt chẽ qua vòng đời:")
    add_body_p(doc, "1. New / Open: Tester phát hiện sai lệch giữa kết quả thực tế và kết quả mong đợi, tạo ticket lỗi với đầy đủ các bước tái hiện (Steps to Reproduce), ảnh chụp màn hình và nhật ký log.")
    add_body_p(doc, "2. Assigned: Trưởng nhóm phân công cho lập trình viên phụ trách module tương ứng để phân tích nguyên nhân gốc rễ.")
    add_body_p(doc, "3. In Progress: Lập trình viên tiến hành sửa mã nguồn, thực thi unit test và kiểm tra cục bộ.")
    add_body_p(doc, "4. Resolved / Fixed: Lập trình viên đẩy mã nguồn đã sửa lên nhánh kiểm thử và bàn giao lại cho Tester.")
    add_body_p(doc, "5. Verified / Re-test: Tester thực hiện kiểm thử lại theo đúng các bước tái hiện. Nếu đạt, ticket được chuyển sang trạng thái Closed. Nếu chưa đạt, chuyển Re-opened.")
    add_body_p(doc, "6. Closed: Lỗi đã được khắc phục hoàn toàn và nghiệm thu vào nhánh phát hành chính thức.")

    # 5. Append Conclusion
    add_heading_1(doc, "PHẦN KẾT LUẬN")
    add_heading_2(doc, "1. Kết quả đạt được")
    add_body_p(doc, "- Xây dựng thành công hệ thống phần mềm Coffee Shop POS hoàn chỉnh, hoạt động ổn định và tích hợp chặt chẽ 6 mẫu thiết kế phần mềm cốt lõi (Factory Method, Decorator, Strategy, Adapter, State, Observer).")
    add_body_p(doc, "- Thực thi kiểm thử hộp trắng chuyên sâu cho các đơn vị mã nguồn xử lý logic phức tạp, đạt 100% độ bao phủ câu lệnh (Statement Coverage) và độ bao phủ nhánh độc lập (Branch Coverage) theo phương pháp McCabe.")
    add_body_p(doc, "- Thiết kế và thực thi thành công bộ 20 Test Cases kiểm thử hộp đen phủ kín 20 chức năng của hệ thống với tỷ lệ Đạt 100%.")
    add_body_p(doc, "- Xây dựng hệ thống kiểm thử tự động E2E với Playwright, tự động hóa quy trình kiểm thử hồi quy và ghi hình video chất lượng cao phục vụ nghiệm thu.")

    add_heading_2(doc, "2. Ưu điểm và nhược điểm")
    add_body_p(doc, "Ưu điểm: Kiến trúc phân tầng rõ ràng, tuân thủ nguyên tắc SOLID, hỗ trợ giao diện cảm ứng và bếp pha chế thời gian thực, quy trình kiểm thử kết hợp toàn diện giữa hộp trắng, hộp đen và kiểm thử tự động.")
    add_body_p(doc, "Nhược điểm: Cổng thanh toán (Momo, VNPay) hiện đang ở mức giả lập kết nối (Adapter Sandbox), chưa hỗ trợ triển khai phân tán nhiều chi nhánh qua mạng Internet.")

    add_heading_2(doc, "3. Hướng phát triển")
    add_body_p(doc, "- Tích hợp Webhook kết nối trực tiếp với cổng thanh toán thực tế của VNPay và Momo.")
    add_body_p(doc, "- Phát triển ứng dụng di động (Mobile App) dành cho khách hàng đặt món trực tuyến và tích lũy điểm thưởng.")
    add_body_p(doc, "- Ứng dụng trí tuệ nhân tạo (AI) trong phân tích dữ liệu bán hàng để dự báo xu hướng tiêu thụ nguyên vật liệu.")

    doc.save(target_file)
    print(f"Successfully generated: {target_file}")

if __name__ == '__main__':
    build_full_report_file('docs/BaoCao_CoffeeShopPOS_DesignPattern.docx')
    build_full_report_file('docs/BaoCao_KiemThu_Nhom15_Final.docx')
    build_full_report_file('docs/Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx')

