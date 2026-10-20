# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document(out_path):
    doc = docx.Document()
    
    # Page setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_bullet_p(text, bold_prefix="", space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        # Bullet symbol •
        r_b = p.add_run("•  ")
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
            
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_subbullet_p(text, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.55)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        # Sub-bullet symbol o
        r_b = p.add_run("o  ")
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(11)
        
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_num_p(num_str, text, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.55)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(num_str + "  ")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(13)
        
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_formula_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.55)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_conclusion_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.55)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run("=> " + text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        return p

    def add_code(code_str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, 'F5F5F5')
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(code_str)
        run.font.name = 'Consolas'
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_pic(img_path, width_in=2.6, caption=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(13)


    # ==========================================================================
    # 4.3. Kiểm thử hàm restockForOrder()
    # ==========================================================================
    add_h2("4.3. Kiểm thử hàm restockForOrder()")
    
    add_h3("4.3.1. Mô tả chức năng")
    add_body_p("Hàm hoàn trả lại nguyên liệu vào kho cho một đơn hàng đã bị trừ kho trước đó (ví dụ khi hủy đơn hàng). Nếu đơn hàng chưa từng bị trừ kho (inventoryDeducted == false), hàm kết thúc ngay lập tức. Ngược lại, hàm tính toán tổng lượng nguyên liệu đơn hàng đã dùng, duyệt qua từng nguyên liệu để cộng hoàn lại vào kho qua kho dữ liệu với mã lý do \"ORDER_RESTOCK\", sau đó đặt lại cờ inventoryDeducted = false để tránh hoàn kho trùng lặp.")

    add_h3("4.3.2. Đoạn mã nguồn")
    code_restock_for_order = """public void restockForOrder(Order order) {
    if (!order.isInventoryDeducted()) { // (1)
        return; // (2)
    }
    Map<String, Double> requirements = calculateRequirements(order); // (3)
    requirements.forEach((name, quantity) -> { // (4)
        InventoryItem item = getItem(name); // (5)
        repository.adjustInventory(item.getId(), quantity, order.getId(), "ORDER_RESTOCK"); // (6)
    });
    order.setInventoryDeducted(false); // (7)
} // (8)"""
    add_code(code_restock_for_order)

    add_h3("4.3.3. Lập đồ thị dòng điều khiển cơ bản")
    add_pic("docs/diagrams/cfg_restock_for_order.png", width_in=2.5, caption="Hình 4.5 Đồ thị dòng điều khiển cơ bản hàm restockForOrder()")

    add_h3("4.3.4. Tính số đường thi hành tuyến tính độc lập.")
    add_bullet_p("Đồ thị bên có 2 nút quyết định nhị phân (1, 4) nên có độ phức tạp:")
    add_formula_p("C = 2 + 1 = 3")
    add_bullet_p("3 đường thi hành tuyến tính lập cơ bản là:")
    add_num_p("1.", "1 → 2 → End", space_after=6)
    add_num_p("2.", "1 → 3 → 4 → 7 → End", space_after=6)
    add_num_p("3.", "1 → 3 → 4 → 5 → 6 → 4 → 7 → End", space_after=6)

    add_h3("4.3.5. Thiết kế test cases cho từng đường")
    add_bullet_p("", bold_prefix="Test case cho đường 1:")
    add_subbullet_p("order.isInventoryDeducted() == false (đơn hàng chưa bị trừ kho).")
    add_subbullet_p("Ví dụ: order có id = 1, inventoryDeducted = false.")
    add_subbullet_p("Kết quả kỳ vọng: Thoát ngay tại lệnh return, không tính toán định lượng nguyên liệu, không gọi repository.adjustInventory(), cờ inventoryDeducted của đơn hàng giữ nguyên giá trị false.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 2:")
    add_subbullet_p("order.isInventoryDeducted() == true.")
    add_subbullet_p("calculateRequirements(order) trả về danh sách rỗng (đơn hàng không có món nào cần trừ kho, requirements = {}).")
    add_subbullet_p("Kết quả kỳ vọng: Vòng lặp forEach tại nút (4) không thực hiện lần nào, không gọi getItem() và adjustInventory(), cập nhật order.setInventoryDeducted(false) thành false và kết thúc bình thường.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 3:")
    add_subbullet_p("order.isInventoryDeducted() == true.")
    add_subbullet_p("calculateRequirements(order) trả về danh sách nguyên liệu hợp lệ (ví dụ: requirements = {\"Cà chua\": 2.0}, trong kho tồn tại item{id=10, name=\"Cà chua\"}).")
    add_subbullet_p("Kết quả kỳ vọng: Vòng lặp forEach duyệt qua nguyên liệu, gọi getItem(\"Cà chua\") tại nút (5), gọi repository.adjustInventory(10, 2.0, order.getId(), \"ORDER_RESTOCK\") tại nút (6) để cộng lại tồn kho, sau đó chuyển sang nút (7) cập nhật inventoryDeducted = false.", space_after=6)

    add_h3("4.3.6. Lập đồ thị dòng dữ liệu chung")
    add_pic("docs/diagrams/dfg_restock_for_order.png", width_in=3.0, caption="Hình 4.6: Đồ thị dòng dữ liệu chung hàm restockForOrder()")

    add_h3("4.3.7. Kiểm thử đời sống của từng biến trong đơn vị mã nguồn")
    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến order:")
    add_num_p("1.", "Kịch bản 1: ~duk")
    add_num_p("2.", "Kịch bản 2: ~duuuk")
    add_num_p("3.", "Kịch bản 3: ~duuuuk")
    add_conclusion_p("Cả 3 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến requirements:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: ~duk")
    add_num_p("3.", "Kịch bản 3: ~duuk")
    add_conclusion_p("Cả 3 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến item:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: Không được khởi tạo.")
    add_num_p("3.", "Kịch bản 3: ~duk")
    add_conclusion_p("Biến chỉ được định nghĩa trong phạm vi thân lambda tại nút (5) và sử dụng tại nút (6). Kịch bản 3 không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến name:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: Không được khởi tạo.")
    add_num_p("3.", "Kịch bản 3: ~duk")
    add_conclusion_p("Biến nhận khóa của Map khi duyệt, dùng làm đối số cho getItem(name). Không có cặp hoạt động bất thường.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến quantity:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: Không được khởi tạo.")
    add_num_p("3.", "Kịch bản 3: ~duk")
    add_conclusion_p("Biến nhận giá trị số lượng của Map, dùng làm đối số cho adjustInventory. Không có cặp hoạt động bất thường.")

    doc.add_page_break()

    # ==========================================================================
    # 4.4. Kiểm thử hàm validateAvailable()
    # ==========================================================================
    add_h2("4.4. Kiểm thử hàm validateAvailable()")
    
    add_h3("4.4.1. Mô tả chức năng")
    add_body_p("Kiểm tra lượng tồn kho hiện tại có đáp ứng đủ cho toàn bộ định mức yêu cầu (requirements) hay không. Hàm duyệt qua từng cặp nguyên liệu: nếu phát hiện bất kỳ nguyên liệu nào có tồn kho nhỏ hơn số lượng cần, hàm lập tức ném ngoại lệ InventoryException kèm thông báo chi tiết (tên món, lượng cần, lượng hiện có và đơn vị tính). Nếu tất cả nguyên liệu đều đủ, phương thức kết thúc bình thường mà không làm thay đổi dữ liệu kho.")

    add_h3("4.4.2. Đoạn mã nguồn")
    code_validate_available = """private void validateAvailable(Map<String, Double> requirements) {
    for (Map.Entry<String, Double> requirement : requirements.entrySet()) { // (1)
        InventoryItem item = getItem(requirement.getKey()); // (2)
        if (item.getQuantity() < requirement.getValue()) { // (3)
            throw new InventoryException("Not enough inventory: " + item.getName()
                    + ". Required " + requirement.getValue() + " " + item.getUnit()
                    + ", available " + item.getQuantity() + " " + item.getUnit() + "."); // (4)
        }
    }
    return; // (5)
}"""
    add_code(code_validate_available)

    add_h3("4.4.3. Lập đồ thị dòng điều khiển cơ bản")
    add_pic("docs/diagrams/cfg_validate_available.png", width_in=2.5, caption="Hình 4.7 Đồ thị dòng điều khiển cơ bản hàm validateAvailable()")

    add_h3("4.4.4. Tính số đường thi hành tuyến tính độc lập.")
    add_bullet_p("Đồ thị bên có 2 nút quyết định nhị phân (1, 3) nên có độ phức tạp:")
    add_formula_p("C = 2 + 1 = 3")
    add_bullet_p("3 đường thi hành tuyến tính lập cơ bản là:")
    add_num_p("1.", "1 → 5 → End", space_after=6)
    add_num_p("2.", "1 → 2 → 3 → 4 → End", space_after=6)
    add_num_p("3.", "1 → 2 → 3 → 1 → 5 → End", space_after=6)

    add_h3("4.4.5. Thiết kế test cases cho từng đường")
    add_bullet_p("", bold_prefix="Test case cho đường 1:")
    add_subbullet_p("requirements rỗng (requirements = {}).")
    add_subbullet_p("Kết quả kỳ vọng: Vòng lặp kết thúc ngay tại nút (1), chuyển sang nút (5) kết thúc hàm bình thường; không gọi getItem(), không ném ngoại lệ.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 2:")
    add_subbullet_p("requirements có ít nhất 1 mục mà lượng tồn kho trong kho không đáp ứng đủ (item.getQuantity() < requirement.getValue()).")
    add_subbullet_p("Ví dụ: requirements = {\"Cà chua\": 5.0}, trong kho chỉ có item \"Cà chua\" với quantity = 3.0, unit = \"kg\".")
    add_subbullet_p("Kết quả kỳ vọng: Điều kiện rẽ nhánh tại nút (3) nhận giá trị True (3.0 < 5.0), ném ngoại lệ InventoryException với thông báo \"Not enough inventory: Cà chua. Required 5.0 kg, available 3.0 kg.\". Phương thức kết thúc sớm bằng ngoại lệ tại nút (4).", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 3:")
    add_subbullet_p("requirements có 1 mục và lượng tồn kho đáp ứng đủ (item.getQuantity() >= requirement.getValue()).")
    add_subbullet_p("Ví dụ: requirements = {\"Cà phê sữa\": 2.0}, trong kho có tồn quantity = 3000.0 (g).")
    add_subbullet_p("Kết quả kỳ vọng: Điều kiện tại nút (3) nhận giá trị False, quay lại nút (1) duyệt hết danh sách, đi qua nút (5) và kết thúc hàm thành công không ném ngoại lệ.", space_after=6)

    add_h3("4.4.6. Lập đồ thị dòng dữ liệu chung")
    add_pic("docs/diagrams/dfg_validate_available.png", width_in=3.0, caption="Hình 4.8: Đồ thị dòng dữ liệu chung hàm validateAvailable()")

    add_h3("4.4.7. Kiểm thử đời sống của từng biến trong đơn vị mã nguồn")
    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến requirements:")
    add_num_p("1.", "Kịch bản 1: ~duk")
    add_num_p("2.", "Kịch bản 2: ~duk (giống kịch bản 1)")
    add_num_p("3.", "Kịch bản 3: ~duuk")
    add_conclusion_p("Cả 3 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến requirement:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: ~duuuk")
    add_num_p("3.", "Kịch bản 3: ~duuk")
    add_conclusion_p("Cả 3 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến item:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: ~duuk")
    add_num_p("3.", "Kịch bản 3: ~duk")
    add_conclusion_p("Cả 3 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    doc.add_page_break()

    # ==========================================================================
    # 4.5. Kiểm thử hàm restockItem()
    # ==========================================================================
    add_h2("4.5. Kiểm thử hàm restockItem()")
    
    add_h3("4.5.1. Mô tả chức năng")
    add_body_p("Hàm cho phép nhân viên nhập thêm một lượng nguyên liệu vào kho theo phương thức thủ công. Hàm kiểm tra số lượng bổ sung amount phải lớn hơn 0; nếu <= 0 sẽ ném ngoại lệ IllegalArgumentException. Khi số lượng hợp lệ, hàm truy vấn nguyên liệu theo id và thực hiện cộng kho thông qua repository.adjustInventory() với lý do giao dịch \"MANUAL_RESTOCK\" và mã đơn hàng là null.")

    add_h3("4.5.2. Đoạn mã nguồn")
    code_restock_item = """public void restockItem(int id, double amount) {
    if (amount <= 0) { // (1)
        throw new IllegalArgumentException("Restock amount must be greater than 0."); // (2)
    }
    InventoryItem item = getInventoryItemById(id); // (3)
    repository.adjustInventory(item.getId(), amount, null, "MANUAL_RESTOCK"); // (4)
    return; // (5)
}"""
    add_code(code_restock_item)

    add_h3("4.5.3. Lập đồ thị dòng điều khiển cơ bản")
    add_pic("docs/diagrams/cfg_restock_item.png", width_in=2.5, caption="Hình 4.9 Đồ thị dòng điều khiển cơ bản hàm restockItem()")

    add_h3("4.5.4. Tính số đường thi hành tuyến tính độc lập.")
    add_bullet_p("Đồ thị bên có 1 nút quyết định nhị phân (1) nên có độ phức tạp:")
    add_formula_p("C = 1 + 1 = 2")
    add_bullet_p("2 đường thi hành tuyến tính lập cơ bản là:")
    add_num_p("1.", "1 → 2 → End", space_after=6)
    add_num_p("2.", "1 → 3 → 4 → 5 → End", space_after=6)

    add_h3("4.5.5. Thiết kế test cases cho từng đường")
    add_bullet_p("", bold_prefix="Test case cho đường 1:")
    add_subbullet_p("amount <= 0.")
    add_subbullet_p("Ví dụ: id = 1, amount = 0 (hoặc amount = -5.0).")
    add_subbullet_p("Kết quả kỳ vọng: Ném ngoại lệ IllegalArgumentException với thông báo \"Restock amount must be greater than 0.\". Không tìm kiếm nguyên liệu và không gọi repository.adjustInventory().", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 2:")
    add_subbullet_p("amount > 0 và id nguyên liệu tồn tại hợp lệ trong kho.")
    add_subbullet_p("Ví dụ: id = 1, amount = 5.5, tìm thấy item{id=1} trong kho.")
    add_subbullet_p("Kết quả kỳ vọng: Tìm thấy nguyên liệu tại nút (3), gọi repository.adjustInventory(1, 5.5, null, \"MANUAL_RESTOCK\") tại nút (4) đúng 1 lần, kết thúc phương thức bình thường tại nút (5).", space_after=6)

    add_h3("4.5.6. Lập đồ thị dòng dữ liệu chung")
    add_pic("docs/diagrams/dfg_restock_item.png", width_in=3.0, caption="Hình 4.10: Đồ thị dòng dữ liệu chung hàm restockItem()")

    add_h3("4.5.7. Kiểm thử đời sống của từng biến trong đơn vị mã nguồn")
    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến id:")
    add_num_p("1.", "Kịch bản 1: ~dk")
    add_num_p("2.", "Kịch bản 2: ~duk")
    add_conclusion_p("Kịch bản 1 xuất hiện chuỗi dk do hàm ném ngoại lệ dừng sớm khi amount <= 0. Đây là trường hợp kết thúc sớm có chủ đích. Kịch bản 2 không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến amount:")
    add_num_p("1.", "Kịch bản 1: ~duk")
    add_num_p("2.", "Kịch bản 2: ~duuk")
    add_conclusion_p("Cả 2 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến item:")
    add_num_p("1.", "Kịch bản 1: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 2: ~duk")
    add_conclusion_p("Biến chỉ được định nghĩa tại nút (3) và sử dụng tại nút (4) để lấy ID. Kịch bản 2 không chứa cặp đôi hoạt động bất thường nào cả.")

    doc.add_page_break()

    # ==========================================================================
    # 4.15. Kiểm thử hàm addRequirements()
    # ==========================================================================
    add_h2("4.15. Kiểm thử hàm addRequirements()")
    
    add_h3("4.15.1. Mô tả chức năng")
    add_body_p("Hàm tính toán tổng hợp định lượng nguyên liệu và vật tư tiêu hao cần dùng cho một dòng sản phẩm (OrderItem) trong đơn hàng và cộng dồn vào bảng requirements. Hàm kiểm tra danh mục công thức định lượng (RecipeItem): nếu không có công thức thì áp dụng mặc định hạt cà phê; nếu có thì duyệt công thức để nhân theo số lượng món. Tiếp đó, hàm kiểm tra các từ khóa topping/tùy biến trong chuỗi mô tả món (description) như trân châu, pudding, kem cheese, extra shot, size L, kem muối, v.v. để cộng dồn định lượng tương ứng vào bảng nguyên liệu.")

    add_h3("4.15.2. Đoạn mã nguồn")
    code_add_requirements = """private void addRequirements(Map<String, Double> requirements, OrderItem item) {
    double quantity = item.getQuantity(); // (1)
    List<RecipeItem> recipeItems = repository.getRecipeItems(item.getBeverageId()); // (2)
    if (recipeItems.isEmpty()) { // (3)
        add(requirements, "Coffee beans", 10 * quantity); // (4)
    } else {
        for (RecipeItem recipeItem : recipeItems) { // (5)
            InventoryItem invItem = getInventoryItemById(recipeItem.getInventoryItemId()); // (6)
            add(requirements, invItem.getName(), recipeItem.getQuantityRequired() * quantity); // (7)
        }
    }
    String description = item.getBeverage().getDescription(); // (8)
    if (hasAny(description, "Tran chau", "Trân châu trắng")) add(requirements, "Pearl", 30 * quantity); // (9), (10)
    if (description.contains("Pudding")) add(requirements, "Pudding", 35 * quantity); // (11), (12)
    if (description.contains("Kem cheese")) add(requirements, "Cream cheese", 25 * quantity); // (13), (14)
    if (description.contains("Extra shot")) add(requirements, "Coffee beans", 8 * quantity); // (15), (16)
    if (description.contains("Size L")) add(requirements, "Cup L", quantity); // (17), (18)
    if (description.contains("Kem muối")) add(requirements, "Salted cream", 25 * quantity); // (19), (20)
    if (description.contains("Thach cafe")) add(requirements, "Coffee jelly", 35 * quantity); // (21), (22)
    if (description.contains("Kem vani")) add(requirements, "Vanilla cream", 25 * quantity); // (23), (24)
    if (hasAny(description, "Duong den", "Trân châu đường đen")) add(requirements, "Brown sugar syrup", 20 * quantity); // (25), (26)
    if (description.contains("Trân châu đường đen")) add(requirements, "Pearl", 30 * quantity); // (27), (28)
    return; // (29)
}"""
    add_code(code_add_requirements)

    add_h3("4.15.3. Lập đồ thị dòng điều khiển cơ bản")
    add_pic("docs/diagrams/cfg_add_requirements.png", width_in=2.6, caption="Hình 4.11 Đồ thị dòng điều khiển cơ bản hàm addRequirements()")

    add_h3("4.15.4. Tính số đường thi hành tuyến tính độc lập.")
    add_bullet_p("Đồ thị có 12 nút quyết định nhị phân: (3), (5), (9), (11), (13), (15), (17), (19), (21), (23), (25), (27) nên có độ phức tạp:")
    add_formula_p("C = 12 + 1 = 13")
    add_bullet_p("13 đường thi hành tuyến tính lập cơ bản là:")
    add_num_p("1.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("2.", "1 → 2 → 3 → 4 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("3.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("4.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 10 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("5.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 12 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("6.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 14 → 15 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("7.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 16 → 17 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("8.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 18 → 19 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("9.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 20 → 21 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("10.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 22 → 23 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("11.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 24 → 25 → 27 → 29 → End", space_after=6)
    add_num_p("12.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 26 → 27 → 29 → End", space_after=6)
    add_num_p("13.", "1 → 2 → 3 → 5 → 6 → 7 → 5 → 8 → 9 → 11 → 13 → 15 → 17 → 19 → 21 → 23 → 25 → 26 → 27 → 28 → 29 → End", space_after=6)

    add_h3("4.15.5. Thiết kế test cases cho từng đường")
    add_bullet_p("", bold_prefix="Test case cho đường 1:")
    add_subbullet_p("Đơn hàng có công thức gồm 1 nguyên liệu, không có phụ liệu hay topping.")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê sữa\", recipeItems = [{inventoryItemId=5, name=\"Milk\", quantityRequired=20}].")
    add_subbullet_p("Kết quả kỳ vọng: requirements được ghi nhận kết quả: {\"Milk\": 40.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 2:")
    add_subbullet_p("Đơn hàng có công thức định lượng rỗng (recipeItems.isEmpty() == true), không có phụ liệu.")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê đen\", recipeItems = [].")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (4), thêm mặc định hạt cà phê: requirements = {\"Coffee beans\": 20.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 3:")
    add_subbullet_p("Đơn hàng có công thức chứa từ 2 nguyên liệu trở lên, vòng lặp (5) chạy 2 lần.")
    add_subbullet_p("Dữ liệu: quantity = 2, recipeItems = [{5, \"Milk\", 20}, {6, \"Sugar\", 10}].")
    add_subbullet_p("Kết quả kỳ vọng: requirements = {\"Milk\": 40.0, \"Sugar\": 20.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 4:")
    add_subbullet_p("Có topping Trân châu (nút 9 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Trà sữa Tran chau\", công thức sữa {\"Milk\": 20}.")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (10), cộng thêm 60g trân châu: requirements = {\"Milk\": 40.0, \"Pearl\": 60.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 5:")
    add_subbullet_p("Có topping Pudding (nút 11 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Trà sữa Pudding\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (12), cộng thêm 70g pudding: requirements = {\"Milk\": 40.0, \"Pudding\": 70.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 6:")
    add_subbullet_p("Có topping Kem cheese (nút 13 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Trà sữa Kem cheese\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (14), requirements = {\"Milk\": 40.0, \"Cream cheese\": 50.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 7:")
    add_subbullet_p("Thêm Extra shot cà phê (nút 15 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê Extra shot\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (16), cộng thêm 16g hạt cà phê: requirements = {\"Milk\": 40.0, \"Coffee beans\": 16.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 8:")
    add_subbullet_p("Chọn ly Size L (nút 17 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Trà sữa Size L\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (18), cộng thêm 2 ly size L: requirements = {\"Milk\": 40.0, \"Cup L\": 2.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 9:")
    add_subbullet_p("Có topping Kem muối (nút 19 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê Kem muối\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (20), requirements = {\"Milk\": 40.0, \"Salted cream\": 50.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 10:")
    add_subbullet_p("Có topping Thạch cafe (nút 21 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê Thach cafe\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (22), requirements = {\"Milk\": 40.0, \"Coffee jelly\": 70.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 11:")
    add_subbullet_p("Có topping Kem vani (nút 23 True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Cà phê Kem vani\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (24), requirements = {\"Milk\": 40.0, \"Vanilla cream\": 50.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 12:")
    add_subbullet_p("Có topping Đường đen (nút 25 True, nút 27 False).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Sữa tươi Duong den\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua nút (26), requirements = {\"Milk\": 40.0, \"Brown sugar syrup\": 40.0}.", space_after=6)

    add_bullet_p("", bold_prefix="Test case cho đường 13:")
    add_subbullet_p("Có Trân châu đường đen (cả nút 25 và 27 cùng True).")
    add_subbullet_p("Dữ liệu: quantity = 2, description = \"Sữa tươi Trân châu đường đen\".")
    add_subbullet_p("Kết quả kỳ vọng: Đi qua cả nút (26) và nút (28), requirements = {\"Milk\": 40.0, \"Brown sugar syrup\": 40.0, \"Pearl\": 60.0}.", space_after=6)

    add_h3("4.15.6. Lập đồ thị dòng dữ liệu chung")
    add_pic("docs/diagrams/dfg_add_requirements.png", width_in=3.0, caption="Hình 4.12: Đồ thị dòng dữ liệu chung hàm addRequirements()")

    add_h3("4.15.7. Kiểm thử đời sống của từng biến trong đơn vị mã nguồn")
    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến requirements:")
    add_num_p("1.", "Kịch bản 1, 2: ~duk")
    add_num_p("2.", "Kịch bản 3, 4, 5, 6, 7, 8, 9, 10, 11, 12: ~duuk")
    add_num_p("3.", "Kịch bản 13: ~duuuk")
    add_conclusion_p("Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến item:")
    add_conclusion_p("Cả 13 kịch bản: ~duuuk")
    add_conclusion_p("Biến được định nghĩa từ tham số truyền vào, dùng ở nút (1), (2), (8) và kết thúc hàm. Không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến quantity:")
    add_num_p("1.", "Kịch bản 1, 2: ~duk")
    add_num_p("2.", "Kịch bản 3, 4, 5, 6, 7, 8, 9, 10, 11, 12: ~duuk")
    add_num_p("3.", "Kịch bản 13: ~duuuk")
    add_conclusion_p("Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến recipeItems:")
    add_num_p("1.", "Kịch bản 2: ~duk")
    add_num_p("2.", "Kịch bản 1, 4..13: ~duuuk")
    add_num_p("3.", "Kịch bản 3: ~duuuuk")
    add_conclusion_p("Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến recipeItem:")
    add_num_p("1.", "Kịch bản 2: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 1, 4..13: ~duuk")
    add_num_p("3.", "Kịch bản 3: ~duuduuk")
    add_conclusion_p("Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến invItem:")
    add_num_p("1.", "Kịch bản 2: Không được khởi tạo.")
    add_num_p("2.", "Kịch bản 1, 4..13: ~duk")
    add_num_p("3.", "Kịch bản 3: ~duduk")
    add_conclusion_p("Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    add_bullet_p("", bold_prefix="Kiểm thử đời sống biến description:")
    add_conclusion_p("Cả 13 kịch bản: ~duuuuuuuuuuk")
    add_conclusion_p("Biến được định nghĩa tại nút (8) và sử dụng tuần tự qua các biểu thức so khớp điều kiện từ nút (9) đến (27). Cả 13 kịch bản đều không chứa cặp đôi hoạt động bất thường nào cả.")

    doc.save(out_path)
    print(f"SUCCESS: Created Word document at {out_path}")

create_document("docs/KiemThuHopTrang_4Ham_ChuanFormat.docx")
