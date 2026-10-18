# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

testcases = [
    {
        'stt': 1,
        'id': 'TC_BB_01',
        'module': 'F01 - Đăng nhập',
        'desc': 'Đăng nhập thành công với tài khoản Thu ngân hợp lệ',
        'pre': 'Hệ thống đang mở tại màn hình LoginView. Tài khoản cashier đã tồn tại.',
        'data': 'Username: cashier\nPassword: 123',
        'steps': '1. Mở màn hình Đăng nhập.\n2. Nhập Username = "cashier".\n3. Nhập Password = "123".\n4. Nhấn nút "Đăng Nhập Vào Ca".',
        'expected': 'Đăng nhập thành công, điều hướng trực tiếp vào giao diện bán hàng POSView với đúng thông tin thu ngân.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Nguyễn Nhật Huy'
    },
    {
        'stt': 2,
        'id': 'TC_BB_02',
        'module': 'F01 - Đăng nhập',
        'desc': 'Đăng nhập thất bại khi nhập sai mật khẩu',
        'pre': 'Ứng dụng đang hiển thị màn hình LoginView.',
        'data': 'Username: cashier\nPassword: sai_mat_khau',
        'steps': '1. Nhập Username = "cashier".\n2. Nhập Password = "sai_mat_khau".\n3. Nhấn nút "Đăng Nhập Vào Ca".',
        'expected': 'Hệ thống từ chối đăng nhập, hiển thị thông báo lỗi "Tên đăng nhập hoặc mật khẩu không chính xác!", giữ nguyên tại LoginView.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Nguyễn Nhật Huy'
    },
    {
        'stt': 3,
        'id': 'TC_BB_03',
        'module': 'F02 - Quản lý menu',
        'desc': 'Quản trị viên thêm món đồ uống mới với thông tin hợp lệ',
        'pre': 'Quản trị viên (admin) đã đăng nhập và đang mở tab "Menu" trong AdminView.',
        'data': 'Tên món: "Trà Vải Hoa Hồng"\nĐơn giá: 38000\nDanh mục: TEA',
        'steps': '1. Tại tab Menu, điền Tên món = "Trà Vải Hoa Hồng".\n2. Nhập Đơn giá = 38000.\n3. Chọn Danh mục = TEA.\n4. Nhấn nút "Thêm Món".',
        'expected': 'Hệ thống lưu món mới vào cơ sở dữ liệu, hiển thị trong danh sách Menu quản trị và xuất hiện ngay trên lưới bán hàng POS.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Trần Hoàng Phúc Quân'
    },
    {
        'stt': 4,
        'id': 'TC_BB_04',
        'module': 'F02 - Quản lý menu',
        'desc': 'Thêm món đồ uống thất bại khi giá âm hoặc để trống tên món',
        'pre': 'Đang ở tab Quản lý Menu trong AdminView.',
        'data': 'Tên món: ""\nĐơn giá: -5000\nDanh mục: COFFEE',
        'steps': '1. Để trống trường Tên món.\n2. Nhập Đơn giá = -5000.\n3. Chọn Danh mục = COFFEE.\n4. Nhấn "Thêm Món".',
        'expected': 'Hệ thống báo lỗi "Name must not be empty" hoặc "Price must be greater than or equal to 0", không lưu bản ghi vào hệ thống.',
        'tech': 'Phân tích giá trị biên (BVA)',
        'status': 'Đạt (Pass)',
        'tester': 'Trần Hoàng Phúc Quân'
    },
    {
        'stt': 5,
        'id': 'TC_BB_05',
        'module': 'F03 - Quản lý Topping',
        'desc': 'Thêm mới loại Topping và cấu hình phụ phí bán lẻ thành công',
        'pre': 'Đang mở tab Topping trong giao diện Quản Trị Trung Tâm.',
        'data': 'Tên Topping: "Thạch Dừa Giòn"\nPhụ phí: 8000',
        'steps': '1. Mở tab Topping.\n2. Nhập Tên Topping = "Thạch Dừa Giòn".\n3. Nhập Phụ phí = 8000.\n4. Nhấn nút "Thêm Topping".',
        'expected': 'Topping mới được tạo với trạng thái active=true, hiển thị trong bảng Topping và sẵn sàng xuất hiện trong modal chọn topping ở POS.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Hoàng Thanh Tú'
    },
    {
        'stt': 6,
        'id': 'TC_BB_06',
        'module': 'F04 - Tạo đơn hàng',
        'desc': 'Chọn món nước từ danh mục và thêm vào giỏ hàng POS',
        'pre': 'Thu ngân đang ở màn hình bán hàng POSView. Giỏ hàng hiện tại đang trống.',
        'data': 'Món: "Cà phê đen"\nSố lượng: 2\nGhi chú: "Ít đường"',
        'steps': '1. Nhấp chọn thẻ món "Cà phê đen" trên lưới menu.\n2. Chọn số lượng = 2.\n3. Nhập ghi chú = "Ít đường".\n4. Nhấn nút "Thêm Vào Đơn".',
        'expected': 'Đơn hàng mới được tạo (trạng thái PENDING), dòng món hiển thị: 2 x Cà phê đen, ghi chú "Ít đường", tạm tính cập nhật chính xác.',
        'tech': 'Kiểm thử theo kịch bản (Use Case)',
        'status': 'Đạt (Pass)',
        'tester': 'Hoàng Thanh Tú'
    },
    {
        'stt': 7,
        'id': 'TC_BB_07',
        'module': 'F04 - Tạo đơn hàng',
        'desc': 'Cập nhật số lượng món về 0 để tự động xóa món khỏi giỏ hàng',
        'pre': 'Đơn hàng PENDING đang có 1 món "Bạc xỉu" số lượng 1.',
        'data': 'itemId = ID của món Bạc xỉu\nSố lượng mới: 0',
        'steps': '1. Tại dòng món Bạc xỉu, nhấn nút giảm số lượng (-) từ 1 xuống 0.\n2. Hoặc gọi hàm updateItemQuantity với quantity = 0.',
        'expected': 'Hệ thống tự động loại bỏ món Bạc xỉu khỏi giỏ hàng, gọi recalculate tính lại tổng tiền = 0 ₫.',
        'tech': 'Phân tích giá trị biên (BVA)',
        'status': 'Đạt (Pass)',
        'tester': 'Phạm Thị Vân Ánh'
    },
    {
        'stt': 8,
        'id': 'TC_BB_08',
        'module': 'F05 - Tính tiền',
        'desc': 'Kiểm tra tính toán tạm tính (Subtotal) khi có nhiều món và số lượng',
        'pre': 'Đơn hàng đang ở trạng thái PENDING.',
        'data': 'Món 1: 2 x Cà phê sữa (30.000₫ = 60.000₫)\nMón 2: 1 x Trà đào (35.000₫ = 35.000₫)',
        'steps': '1. Thêm 2 ly Cà phê sữa vào đơn.\n2. Thêm 1 ly Trà đào vào đơn.\n3. Quan sát phần tổng tiền trên giao diện POS.',
        'expected': 'Tạm tính (Subtotal) = 60.000 + 35.000 = 95.000 ₫. Tổng thanh toán hiển thị chính xác 95.000 ₫ (khi chưa áp dụng giảm giá).',
        'tech': 'Bảng quyết định (Decision Table)',
        'status': 'Đạt (Pass)',
        'tester': 'Phạm Thị Vân Ánh'
    },
    {
        'stt': 9,
        'id': 'TC_BB_09',
        'module': 'F06 - Khuyến mãi',
        'desc': 'Áp dụng chiến lược giảm giá phần trăm (PercentDiscountStrategy 10%)',
        'pre': 'Đơn hàng PENDING có Subtotal = 100.000 ₫.',
        'data': 'Chiến lược giảm giá: PERCENT_10 (Giảm 10%)',
        'steps': '1. Mở danh sách khuyến mãi tại giỏ hàng POS.\n2. Chọn chiến lược "Giảm 10%".\n3. Nhấn Áp dụng khuyến mãi.',
        'expected': 'Tiền giảm giá DiscountAmount = 10.000 ₫. Tổng tiền thanh toán FinalTotal = 100.000 - 10.000 = 90.000 ₫.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Ô Duy Hoàng Thiện'
    },
    {
        'stt': 10,
        'id': 'TC_BB_10',
        'module': 'F06 - Khuyến mãi',
        'desc': 'Áp dụng khuyến mãi Mua 1 Tặng 1 (BuyOneGetOneStrategy) với >= 2 món',
        'pre': 'Đơn hàng PENDING gồm 2 món nước có giá khác nhau.',
        'data': 'Món A: Cà phê đen (25.000₫)\nMón B: Matcha Latte (45.000₫)',
        'steps': '1. Thêm Món A (25.000₫) và Món B (45.000₫) vào giỏ hàng.\n2. Chọn chiến lược "BUY_ONE_GET_ONE".\n3. Hệ thống tự động tính toán lại.',
        'expected': 'Hệ thống tự động miễn phí món có giá thấp nhất (Món A: 25.000₫). Tiền giảm = 25.000 ₫. Tổng thanh toán = 45.000 ₫.',
        'tech': 'Bảng quyết định (Decision Table)',
        'status': 'Đạt (Pass)',
        'tester': 'Ô Duy Hoàng Thiện'
    },
    {
        'stt': 11,
        'id': 'TC_BB_11',
        'module': 'F07 - Cập nhật trạng thái',
        'desc': 'Chuyển trạng thái đơn hàng tuần tự theo State Pattern hợp lệ',
        'pre': 'Đơn hàng vừa tạo thành công ở trạng thái PENDING.',
        'data': 'Mã đơn: Order #1\nLuồng: PENDING -> PREPARING -> READY -> PAID',
        'steps': '1. Thu ngân nhấn "Gửi Sang Bếp" (chuyển PREPARING).\n2. Nhân viên bếp nhấn "Hoàn Tất Món" (chuyển READY).\n3. Thu ngân thực hiện thanh toán (chuyển PAID).',
        'expected': 'Đơn hàng chuyển trạng thái tuần tự và thành công qua từng bước. Màn hình bếp và POS đồng bộ thời gian thực qua Observer Pattern.',
        'tech': 'Kiểm thử chuyển trạng thái (State Transition)',
        'status': 'Đạt (Pass)',
        'tester': 'Nguyễn Nhật Huy'
    },
    {
        'stt': 12,
        'id': 'TC_BB_12',
        'module': 'F07 - Cập nhật trạng thái',
        'desc': 'Chặn thao tác hủy đơn hàng khi đơn đã thanh toán thành công (PAID)',
        'pre': 'Đơn hàng đã được thanh toán và đang ở trạng thái PAID.',
        'data': 'Order #2 (Status = PAID)\nHành động: Hủy đơn (Cancel)',
        'steps': '1. Chọn đơn hàng đã PAID trong danh sách quản lý đơn.\n2. Cố gắng thực hiện thao tác Hủy đơn (gọi hàm cancel).',
        'expected': 'Hệ thống ném InvalidStateTransitionException chặn lại, thông báo "Cannot cancel when order is PAID", trạng thái giữ nguyên.',
        'tech': 'Kiểm thử chuyển trạng thái (State Transition)',
        'status': 'Đạt (Pass)',
        'tester': 'Nguyễn Nhật Huy'
    },
    {
        'stt': 13,
        'id': 'TC_BB_13',
        'module': 'F08 - Thanh toán',
        'desc': 'Thanh toán thành công qua cổng thanh toán Momo (MomoAdapter)',
        'pre': 'Đơn hàng ở trạng thái READY, tổng thanh toán là 60.000 ₫.',
        'data': 'Phương thức thanh toán: "MOMO"\nSố tiền: 60.000 ₫',
        'steps': '1. Nhấn nút "Thanh Toán" tại đơn hàng.\n2. Chọn hình thức "Ví Momo".\n3. Quét mã QR giả lập và xác nhận thanh toán.',
        'expected': 'Cổng MomoAdapter xử lý thành công, sinh mã giao dịch dạng MOMO-XXXXXXXX, tạo bản ghi Payment SUCCESS, chuyển đơn sang PAID.',
        'tech': 'Kiểm thử theo kịch bản (Use Case)',
        'status': 'Đạt (Pass)',
        'tester': 'Trần Hoàng Phúc Quân'
    },
    {
        'stt': 14,
        'id': 'TC_BB_14',
        'module': 'F08 - Thanh toán',
        'desc': 'Xử lý khi cổng thanh toán VNPay trả về kết quả giao dịch thất bại',
        'pre': 'Đơn hàng đang sẵn sàng thanh toán (READY).',
        'data': 'Phương thức: "VNPAY"\nCổng trả về: success = false',
        'steps': '1. Chọn phương thức thanh toán VNPay.\n2. Cổng VNPay mô phỏng giao dịch thất bại (ví dụ: số dư không đủ).',
        'expected': 'Hệ thống thông báo "Thanh toán thất bại, vui lòng thử lại!", không sinh bản ghi Payment, đơn hàng giữ nguyên ở trạng thái READY.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Trần Hoàng Phúc Quân'
    },
    {
        'stt': 15,
        'id': 'TC_BB_15',
        'module': 'F09 & F18 - Hóa đơn',
        'desc': 'Xem trước hóa đơn và xuất tệp hóa đơn điện tử (.txt E-Receipt)',
        'pre': 'Đơn hàng vừa thanh toán thành công (PAID).',
        'data': 'Order #1 (Chi tiết món, phụ phí, giảm giá, phương thức)',
        'steps': '1. Sau khi thanh toán, modal ReceiptPreview hiển thị.\n2. Kiểm tra thông tin hóa đơn.\n3. Nhấn nút "Tải E-Receipt (.txt)".',
        'expected': 'Hiển thị đầy đủ thông tin cửa hàng, danh sách món, tổng tiền. Tệp receipt_order_1.txt tự động tải xuống máy tính với nội dung chuẩn.',
        'tech': 'Kiểm thử theo kịch bản (Use Case)',
        'status': 'Đạt (Pass)',
        'tester': 'Hoàng Thanh Tú'
    },
    {
        'stt': 16,
        'id': 'TC_BB_16',
        'module': 'F10 - Báo cáo',
        'desc': 'Quản trị viên xem báo cáo doanh thu ca làm việc và top món bán chạy',
        'pre': 'Hệ thống đã có nhiều đơn hàng hoàn thành (PAID) trong ngày.',
        'data': 'Tài khoản admin đang đăng nhập.',
        'steps': '1. Mở tab "Báo Cáo & Doanh Thu" trong AdminView.\n2. Xem các chỉ số: Tổng doanh thu, Số đơn hoàn thành, Biểu đồ món bán chạy.',
        'expected': 'Hệ thống thống kê chính xác tổng số tiền bán được từ các đơn PAID, hiển thị danh sách Top Selling Items với số lượng chính xác.',
        'tech': 'Kiểm thử theo kịch bản (Use Case)',
        'status': 'Đạt (Pass)',
        'tester': 'Hoàng Thanh Tú'
    },
    {
        'stt': 17,
        'id': 'TC_BB_17',
        'module': 'F11 - Đăng ký tài khoản',
        'desc': 'Đăng ký tài khoản nhân viên mới hợp lệ và phân quyền đúng vai trò',
        'pre': 'Đang mở tab Đăng Ký Mới trên màn hình LoginView.',
        'data': 'Họ tên: "Lê Văn C"\nUsername: "kitchen02"\nRole: "KITCHEN"\nMật khẩu: "123"\nXác nhận mật khẩu: "123"',
        'steps': '1. Nhập đầy đủ Họ tên, Username, Role và Mật khẩu.\n2. Nhấn nút "Xác Nhận Đăng Ký Tài Khoản".',
        'expected': 'Hệ thống báo "Đăng ký tài khoản thành công!", lưu user mới vào SQLite, tự động chuyển về form đăng nhập với username kitchen02.',
        'tech': 'Phân vùng tương đương (EP)',
        'status': 'Đạt (Pass)',
        'tester': 'Phạm Thị Vân Ánh'
    },
    {
        'stt': 18,
        'id': 'TC_BB_18',
        'module': 'F12 - Đổi mật khẩu',
        'desc': 'Đổi mật khẩu cá nhân thành công khi mật khẩu hiện tại chính xác',
        'pre': 'Nhân viên thu ngân đang đăng nhập trong ca trực.',
        'data': 'Mật khẩu cũ: "123"\nMật khẩu mới: "pos2026"\nXác nhận mật khẩu: "pos2026"',
        'steps': '1. Nhấp biểu tượng chìa khóa đổi mật khẩu trên Navbar.\n2. Nhập mật khẩu cũ = "123".\n3. Nhập mật khẩu mới = "pos2026" và xác nhận.\n4. Nhấn "Lưu Mật Khẩu".',
        'expected': 'Hệ thống cập nhật mật khẩu mới, thông báo thành công và tự động đóng modal. Lần đăng nhập sau sử dụng mật khẩu mới pos2026.',
        'tech': 'Phân tích giá trị biên (BVA)',
        'status': 'Đạt (Pass)',
        'tester': 'Phạm Thị Vân Ánh'
    },
    {
        'stt': 19,
        'id': 'TC_BB_19',
        'module': 'F16 - Tùy biến món Decorator',
        'desc': 'Tùy biến món nước chuyên sâu: Size L, đường, đá và Topping',
        'pre': 'Thu ngân đang chọn món "Cà phê sữa" (giá gốc 30.000 ₫).',
        'data': 'Kích cỡ: Size L (+10.000₫)\nLượng đường: 50%\nLượng đá: 50%\nTopping: "Kem cheese" (+12.000₫)',
        'steps': '1. Nhấp chọn thẻ món Cà phê sữa để mở Decorator Customizer.\n2. Chọn Size L, đường 50%, đá 50%.\n3. Tích chọn Topping "Kem cheese".\n4. Nhấn "+ Thêm Vào Đơn".',
        'expected': 'Sản phẩm được bọc Decorator hoàn chỉnh. Đơn giá món = 30.000 + 10.000 + 12.000 = 52.000 ₫. Mô tả hiển thị đầy đủ tùy chọn.',
        'tech': 'Bảng quyết định (Decision Table)',
        'status': 'Đạt (Pass)',
        'tester': 'Ô Duy Hoàng Thiện'
    },
    {
        'stt': 20,
        'id': 'TC_BB_20',
        'module': 'F17 & F19 - Trừ kho & Cảnh báo',
        'desc': 'Khấu trừ nguyên liệu kho tự động khi chuyển pha chế và phát cảnh báo',
        'pre': 'Kho có 20g Cà phê hạt. Món Cà phê đen yêu cầu 15g Cà phê hạt. Mức cảnh báo kho reorderLevel = 10g.',
        'data': 'Đơn hàng gồm 1 ly Cà phê đen (tiêu hao 15g Cà phê hạt)',
        'steps': '1. Tạo đơn 1 ly Cà phê đen.\n2. Nhấn "Gửi Sang Bếp" (chuyển sang PREPARING).\n3. Hệ thống gọi deductForOrder kiểm tra và trừ kho.',
        'expected': 'Tồn kho Cà phê hạt giảm từ 20g xuống 5g (20 - 15 = 5g). Vì 5g < reorderLevel (10g), hệ thống tự động phát cảnh báo thiếu kho lên chuông thông báo.',
        'tech': 'Phân tích giá trị biên (BVA) + Observer',
        'status': 'Đạt (Pass)',
        'tester': 'Ô Duy Hoàng Thiện'
    }
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'BlackBox_TestCases'

ws.views.sheetView[0].showGridLines = True

# Title banner
ws.merge_cells('A1:K1')
ws['A1'] = 'BẢNG KỊCH BẢN KIỂM THỬ HỘP ĐEN (BLACK-BOX TEST CASES) - COFFEE SHOP POS'
ws['A1'].font = Font(name='Arial', size=15, bold=True, color='FFFFFF')
ws['A1'].fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 38

# Subtitle
ws.merge_cells('A2:K2')
ws['A2'] = 'Đề tài: Hệ Thống Coffee Shop POS | Học kỳ I (2026-2027) | GVHD: Ths. Nguyễn Trần Thi Văn | SVTH: Nhóm 15'
ws['A2'].font = Font(name='Arial', size=10, italic=True, color='262626')
ws['A2'].fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 24

headers = [
    'STT', 'Mã Test Case', 'Chức năng / Module', 'Mục đích kiểm thử',
    'Điều kiện tiền đề', 'Dữ liệu đầu vào (Test Data)', 'Các bước thực hiện (Test Steps)',
    'Kết quả mong đợi (Expected Result)', 'Kỹ thuật kiểm thử', 'Trạng thái', 'Người thực hiện'
]

header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
thin_border = Border(
    left=Side(style='thin', color='B0B0B0'),
    right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='B0B0B0'),
    bottom=Side(style='thin', color='B0B0B0')
)

ws.row_dimensions[3].height = 28
for col_num, h in enumerate(headers, 1):
    c = ws.cell(row=3, column=col_num)
    c.value = h
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = thin_border

alt_fill = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
pass_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
pass_font = Font(name='Arial', size=9, bold=True, color='375623')
regular_font = Font(name='Arial', size=9, color='000000')
code_font = Font(name='Arial', size=9, bold=True, color='1F4E79')

for idx, tc in enumerate(testcases, start=4):
    row_fill = alt_fill if idx % 2 == 0 else white_fill
    ws.row_dimensions[idx].height = 55
    
    # STT
    c = ws.cell(row=idx, column=1, value=tc['stt'])
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # ID
    c = ws.cell(row=idx, column=2, value=tc['id'])
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.font = code_font
    c.fill = row_fill
    c.border = thin_border
    
    # Module
    c = ws.cell(row=idx, column=3, value=tc['module'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = Font(name='Arial', size=9, bold=True, color='262626')
    c.fill = row_fill
    c.border = thin_border
    
    # Desc
    c = ws.cell(row=idx, column=4, value=tc['desc'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Pre
    c = ws.cell(row=idx, column=5, value=tc['pre'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Data
    c = ws.cell(row=idx, column=6, value=tc['data'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Steps
    c = ws.cell(row=idx, column=7, value=tc['steps'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Expected
    c = ws.cell(row=idx, column=8, value=tc['expected'])
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Technique
    c = ws.cell(row=idx, column=9, value=tc['tech'])
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border
    
    # Status
    c = ws.cell(row=idx, column=10, value=tc['status'])
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.font = pass_font
    c.fill = pass_fill
    c.border = thin_border
    
    # Tester
    c = ws.cell(row=idx, column=11, value=tc['tester'])
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.font = regular_font
    c.fill = row_fill
    c.border = thin_border

col_widths = {
    1: 6,   # STT
    2: 13,  # ID
    3: 20,  # Module
    4: 28,  # Mục đích
    5: 24,  # Tiền đề
    6: 24,  # Dữ liệu
    7: 30,  # Các bước
    8: 32,  # Kết quả mong đợi
    9: 22,  # Kỹ thuật
    10: 13, # Trạng thái
    11: 20  # Người test
}

for col_num, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_num)].width = width

output_path = 'docs/Bang_TestCases_HopDen_CoffeeShopPOS_20TC.xlsx'
wb.save(output_path)
print('File saved successfully to:', output_path)
