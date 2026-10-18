# -*- coding: utf-8 -*-
"""
Tạo file Word (.docx) chuẩn học thuật cho mục:
1.2. TRIỂN KHAI KIỂM THỬ TỰ ĐỘNG CHO PHÂN HỆ BACKEND
Định dạng: Times New Roman, Căn lề chuẩn 2cm, Bảng biểu khoa học, Văn phong học thuật chuẩn mực.
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
import docx
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "docs" / "MUC_1_2_TRIEN_KHAI_KIEM_THU_TU_DONG_BACKEND.docx"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_custom_table(table, col_widths, headers, data_rows, hdr_bg='1F4E79', hdr_text_white=True, font_size_pt=9.5):
    total_dxa = int(sum(col_widths) * 1440)
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
    
    grid_cols = ''.join(f'<w:gridCol {nsdecls("w")} w:w="{int(w * 1440)}"/>' for w in col_widths)
    tblGrid = parse_xml(f'<w:tblGrid {nsdecls("w")}>{grid_cols}</w:tblGrid>')
    table._tbl.insert(table._tbl.index(tblPr) + 1, tblGrid)
    
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for col_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.text = h_text
        set_cell_background(cell, hdr_bg)
        set_cell_margins(cell, top=90, bottom=90, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(font_size_pt)
            r.font.bold = True
            if hdr_text_white:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    for row_idx, r_data in enumerate(data_rows):
        trPr_row = table.rows[row_idx + 1]._tr.get_or_add_trPr()
        trPr_row.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        row_bg = 'F2F4F7' if row_idx % 2 == 1 else 'FFFFFF'
        for col_idx, val in enumerate(r_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(val)
            set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            if col_idx in [0, 1, 5] or (len(str(val)) <= 10 and '\n' not in str(val)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(font_size_pt)
                if 'Pass' in str(val) or 'Đạt' in str(val):
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x80, 0x00)

def add_p(doc, text="", bold_prefix="", indent=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
    else:
        p.paragraph_format.left_indent = Inches(0)
        
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.font.bold = True
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        if italic:
            r.font.italic = True
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x47, 0x61)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return p

def build_docx():
    doc = docx.Document()
    
    # Căn lề chuẩn 2cm
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2)
        s.right_margin = Cm(2)
        
    # Tiêu đề mục báo cáo
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    r_title = p_title.add_run("1.2. TRIỂN KHAI KIỂM THỬ TỰ ĐỘNG CHO PHÂN HỆ BACKEND\n(UNIT TESTING & INTEGRATION TESTING)")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(14.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x47, 0x61)
    
    # -------------------------------------------------------------------------
    # 1.2.1. Mục tiêu và Giải pháp kiểm thử
    # -------------------------------------------------------------------------
    add_h1(doc, "1.2.1. Mục tiêu và Giải pháp kiểm thử tự động Backend")
    add_p(doc, text="Phân hệ Backend đóng vai trò cốt lõi trong việc xử lý logic nghiệp vụ, quản lý trạng thái đơn hàng, tính toán tài chính và lưu trữ dữ liệu của hệ thống Coffee Shop POS. Nhằm đảm bảo tính chính xác của các thuật toán nghiệp vụ và phục vụ kiểm thử hồi quy liên tục, nhóm nghiên cứu đã xây dựng bộ kiểm thử tự động nội bộ bằng ngôn ngữ Java 17.")
    add_p(doc, text="Bộ kiểm thử được phân tách rành mạch thành hai cấp độ chuyên biệt:")
    add_p(doc, text="Kiểm thử đơn vị (Unit Testing): Gồm 18 ca kiểm thử độc lập, sử dụng cơ chế lưu trữ giả lập trên bộ nhớ RAM (InMemoryRepository) để cô lập logic nghiệp vụ của các lớp Service, không phụ thuộc vào hạ tầng cơ sở dữ liệu vật lý.", bold_prefix="• ", indent=True)
    add_p(doc, text="Kiểm thử tích hợp (Integration Testing): Gồm 6 ca kiểm thử, kết nối trực tiếp với hệ quản trị cơ sở dữ liệu SQLite vật lý thông qua các tệp tin tạm thời (tự động khởi tạo và giải phóng sau khi kiểm thử) nhằm đánh giá tính toàn vẹn dữ liệu và nhật ký giao dịch.", bold_prefix="• ", indent=True)
    add_p(doc, text="Toàn bộ 24 ca kiểm thử được điều phối tập trung thông qua TestRunner.java và có thể kích hoạt linh hoạt qua các tập lệnh thực thi tự động (test-unit.bat, test-integration.bat, test.bat).")

    # -------------------------------------------------------------------------
    # 1.2.2. Nội dung các ca kiểm thử
    # -------------------------------------------------------------------------
    add_h1(doc, "1.2.2. Nội dung các chức năng nghiệp vụ đưa vào kiểm thử")
    add_p(doc, text="Nhóm đã thiết lập 24 ca kiểm thử tập trung vào các chức năng nghiệp vụ trọng yếu của hệ thống:")
    
    add_h2(doc, "A. Kiểm thử đơn vị (Unit Testing - 18 ca kiểm thử)")
    add_p(doc, text="TC01 (Xác thực người dùng): Kiểm tra tính đúng đắn của quá trình xác thực tài khoản và phân quyền vai trò Thu ngân (Cashier).", bold_prefix="• ", indent=True)
    add_p(doc, text="TC02 (Tính giá đồ uống tùy biến): Kiểm tra công thức tính lũy tiến đơn giá khi thêm các loại topping và thay đổi kích cỡ đồ uống.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC03 (Tính chiết khấu đơn hàng): Kiểm tra việc áp dụng tỷ lệ giảm giá theo chính sách khuyến mãi trên tổng tiền đơn hàng.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC04 (Vòng đời đơn hàng hợp lệ): Kiểm tra trình tự luân chuyển trạng thái chuẩn của đơn hàng từ Chờ xử lý sang Đang pha chế và Sẵn sàng phục vụ.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC05 (Ràng buộc khóa trạng thái): Kiểm tra cơ chế ngăn chặn can thiệp, không cho phép tiếp tục pha chế khi đơn hàng đã hoàn tất thanh toán.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC06 (Thanh toán qua ví điện tử): Kiểm tra quy trình xử lý giao dịch thành công qua cổng thanh toán điện tử và sinh mã giao dịch.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC07 (Phát thông báo hoàn thành món): Kiểm tra cơ chế gửi thông báo thời gian thực từ quầy pha chế sang màn hình thu ngân khi đơn hàng đã sẵn sàng.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC08 (Quản lý danh mục món nước): Kiểm tra các thao tác thêm mới, cập nhật đơn giá và ngừng phục vụ món trong thực đơn.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC09 (Xác thực dữ liệu thực đơn): Kiểm tra bộ lọc dữ liệu đầu vào, từ chối các trường hợp tên món để trống, đơn giá âm hoặc sai nhóm danh mục.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC10 (Tính toán trừ và hoàn kho nguyên liệu): Kiểm tra việc tự động trừ nguyên liệu theo định mức khi vào bếp và hoàn trả lại kho khi hủy đơn.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC11 (Kiểm soát định mức tồn kho): Kiểm tra điều kiện chặn nhận đơn pha chế khi số lượng tồn kho nguyên liệu không đáp ứng đủ định mức.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC12 (Định dạng nội dung hóa đơn): Kiểm tra việc tổng hợp đầy đủ các thông tin bắt buộc trên hóa đơn thanh toán (tên quán, món, thành tiền, mã giao dịch).", bold_prefix="• ", indent=True)
    add_p(doc, text="TC13 (Kết xuất hóa đơn dạng tệp ảnh): Kiểm tra chức năng tạo tệp hình ảnh hóa đơn hoàn chỉnh phục vụ in ấn.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC14 (Quản lý tài khoản nhân viên): Kiểm tra nghiệp vụ tạo mới tài khoản, khóa quyền truy cập và ngăn chặn tạo trùng tên đăng nhập.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC15 (Xử lý giỏ hàng bán hàng): Kiểm tra việc tự động gộp số lượng khi chọn món trùng lặp và chặn sửa giỏ hàng khi đơn đã gửi chế biến.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC16 (Khởi tạo danh mục theo nhóm sản phẩm): Kiểm tra tính chính xác khi phân loại và khởi tạo các mặt hàng thuộc nhóm Cà phê, Trà và Matcha.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC17 (Quản lý phiên bản cấu hình và kết nối): Kiểm tra tính duy nhất của cấu hình ứng dụng và đối tượng kết nối cơ sở dữ liệu.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC18 (Xử lý sự cố giao dịch thanh toán): Kiểm tra khả năng xử lý ngoại lệ và bảo toàn trạng thái đơn hàng khi cổng thanh toán đối tác trả về lỗi.", bold_prefix="• ", indent=True)

    add_h2(doc, "B. Kiểm thử tích hợp (Integration Testing - 6 ca kiểm thử)")
    add_p(doc, text="TC19 (Lưu trữ và phục hồi dữ liệu đơn hàng): Kiểm tra việc ghi nhận đơn hàng có nhiều tùy biến vào cơ sở dữ liệu SQLite và nạp lại chính xác từ kết nối độc lập.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC20 (Đồng bộ thao tác quản trị vào CSDL): Kiểm tra tính bền vững của các thao tác thêm món, cập nhật topping và khóa nhân viên trong CSDL SQLite.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC21 (Ghi nhận nhật ký kiểm toán - Audit Trail): Kiểm tra việc tự động ghi nhận lịch sử thay đổi trạng thái đơn và nhật ký xuất nhập kho vào các bảng kiểm toán trong CSDL.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC22 (Truy vấn báo cáo vận hành): Kiểm tra tính chính xác của các câu lệnh truy vấn báo cáo lịch sử đơn hàng và biến động kho từ cơ sở dữ liệu.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC23 (Sao lưu và phục hồi cơ sở dữ liệu): Kiểm tra quy trình trích xuất bản sao lưu snapshot tệp CSDL (.db) và khôi phục để kiểm tra tính nhất quán.", bold_prefix="• ", indent=True)
    add_p(doc, text="TC24 (Quản lý định mức công thức pha chế): Kiểm tra việc lưu trữ và xóa bỏ cấu hình định mức nguyên liệu tiêu hao của đồ uống trong CSDL SQLite.", bold_prefix="• ", indent=True)

    # Bảng 1.1: Danh mục 24 Test Cases
    p_cap1 = doc.add_paragraph()
    p_cap1.paragraph_format.space_before = Pt(8)
    p_cap1.paragraph_format.space_after = Pt(4)
    r_cap1 = p_cap1.add_run("Bảng 1.1: Danh mục tổng hợp 24 ca kiểm thử tự động hóa phân hệ Backend")
    r_cap1.font.name = 'Times New Roman'
    r_cap1.font.size = Pt(11)
    r_cap1.font.bold = True
    r_cap1.font.italic = True
    
    tc_headers = ["STT", "Mã TC", "Mức kiểm thử", "Chức năng kiểm thử", "Mục đích kiểm thử", "Kết quả"]
    tc_data = [
        ["1", "TC01", "Unit Test", "Xác thực tài khoản", "Kiểm tra đăng nhập đúng tài khoản và phân quyền vai trò Thu ngân", "Pass"],
        ["2", "TC02", "Unit Test", "Tính giá đồ uống", "Kiểm tra tính lũy tiến đơn giá khi thêm topping và đổi kích cỡ size", "Pass"],
        ["3", "TC03", "Unit Test", "Chiết khấu khuyến mãi", "Kiểm tra tính toán giảm giá theo tỷ lệ phần trăm trên đơn hàng", "Pass"],
        ["4", "TC04", "Unit Test", "Vòng đời đơn hàng", "Kiểm tra luân chuyển trạng thái: Chờ xử lý -> Đang pha chế -> Sẵn sàng", "Pass"],
        ["5", "TC05", "Unit Test", "Khóa trạng thái đơn", "Chặn can thiệp, không cho phép pha chế lại đơn hàng đã thanh toán", "Pass"],
        ["6", "TC06", "Unit Test", "Thanh toán ví điện tử", "Xử lý giao dịch thanh toán thành công và tạo mã giao dịch", "Pass"],
        ["7", "TC07", "Unit Test", "Thông báo pha chế", "Phát tín hiệu thời gian thực từ bếp sang thu ngân khi hoàn tất món", "Pass"],
        ["8", "TC08", "Unit Test", "Quản lý thực đơn", "Thực hiện thao tác thêm món, cập nhật giá và tắt món trong thực đơn", "Pass"],
        ["9", "TC09", "Unit Test", "Xác thực dữ liệu món", "Chặn dữ liệu không hợp lệ: tên rỗng, đơn giá âm, sai danh mục", "Pass"],
        ["10", "TC10", "Unit Test", "Xuất nhập tồn kho", "Tự động trừ kho khi vào chế biến và hoàn trả lại kho khi hủy đơn", "Pass"],
        ["11", "TC11", "Unit Test", "Kiểm soát kho tối thiểu", "Chặn nhận đơn pha chế khi nguyên liệu tồn kho không đủ định mức", "Pass"],
        ["12", "TC12", "Unit Test", "Định dạng hóa đơn", "Kiểm tra đầy đủ các trường thông tin quy định trên hóa đơn thanh toán", "Pass"],
        ["13", "TC13", "Unit Test", "Kết xuất ảnh hóa đơn", "Tạo tệp hình ảnh hóa đơn hoàn chỉnh phục vụ in ấn", "Pass"],
        ["14", "TC14", "Unit Test", "Quản lý nhân viên", "Thêm nhân viên mới, khóa tài khoản và chặn trùng tên đăng nhập", "Pass"],
        ["15", "TC15", "Unit Test", "Xử lý giỏ hàng", "Tự gộp món trùng lặp và chặn sửa giỏ khi đơn đã gửi vào pha chế", "Pass"],
        ["16", "TC16", "Unit Test", "Khởi tạo danh mục món", "Khởi tạo chính xác đồ uống theo từng nhóm Cà phê, Trà, Matcha", "Pass"],
        ["17", "TC17", "Unit Test", "Quản lý phiên bản kết nối", "Bảo đảm tính duy nhất của cấu hình ứng dụng và kết nối CSDL", "Pass"],
        ["18", "TC18", "Unit Test", "Xử lý sự cố thanh toán", "Bảo toàn trạng thái đơn hàng khi cổng thanh toán đối tác trả về lỗi", "Pass"],
        ["19", "TC19", "Integ Test", "Lưu trữ đơn hàng SQLite", "Lưu trữ và phục hồi toàn vẹn đơn hàng có tùy biến trong CSDL SQLite", "Pass"],
        ["20", "TC20", "Integ Test", "Đồng bộ CRUD với CSDL", "Đồng bộ thao tác thêm món, đổi topping và khóa nhân viên vào CSDL", "Pass"],
        ["21", "TC21", "Integ Test", "Nhật ký kiểm toán CSDL", "Ghi nhận lịch sử trạng thái đơn và nhật ký trừ kho vào bảng kiểm toán", "Pass"],
        ["22", "TC22", "Integ Test", "Truy vấn báo cáo vận hành", "Truy vấn chính xác lịch sử đơn hàng và dữ liệu vận hành từ CSDL", "Pass"],
        ["23", "TC23", "Integ Test", "Sao lưu và phục hồi CSDL", "Xuất tệp sao lưu snapshot CSDL và phục hồi kiểm tra tính nhất quán", "Pass"],
        ["24", "TC24", "Integ Test", "Định mức công thức pha chế", "Lưu trữ và xóa bỏ cấu hình định mức nguyên liệu tiêu hao trong CSDL", "Pass"]
    ]
    t1 = doc.add_table(rows=len(tc_data) + 1, cols=len(tc_headers))
    t1.style = 'Table Grid'
    format_custom_table(t1, [0.45, 0.65, 0.95, 1.6, 2.75, 0.6], tc_headers, tc_data, hdr_bg='1F4E79', hdr_text_white=True, font_size_pt=8.5)

    # -------------------------------------------------------------------------
    # 1.2.3. Đo lường Code Coverage bằng JaCoCo
    # -------------------------------------------------------------------------
    add_h1(doc, "1.2.3. Đo lường Độ bao phủ mã nguồn (Code Coverage) bằng công cụ JaCoCo")
    add_p(doc, text="Để đánh giá định lượng độ bao phủ của bộ kiểm thử tự động, công cụ JaCoCo (Java Code Coverage) phiên bản 0.8.14 được tích hợp vào quy trình thực thi. Dữ liệu thực thi được máy ảo JVM ghi nhận vào tệp nhị phân jacoco.exec và xuất báo cáo kết quả chi tiết.")
    
    # Bảng 1.2: Bảng số liệu tổng thể
    p_cap2 = doc.add_paragraph()
    p_cap2.paragraph_format.space_before = Pt(8)
    p_cap2.paragraph_format.space_after = Pt(4)
    r_cap2 = p_cap2.add_run("Bảng 1.2: Bảng tổng hợp mức độ bao phủ mã nguồn Backend đo bằng JaCoCo")
    r_cap2.font.name = 'Times New Roman'
    r_cap2.font.size = Pt(11)
    r_cap2.font.bold = True
    r_cap2.font.italic = True
    
    cov_sum_headers = ["Phân vùng mã nguồn kiểm thử", "Bao phủ dòng lệnh (Line)", "Bao phủ rẽ nhánh (Branch)", "Bao phủ mã máy (Instruction)", "Kết luận"]
    cov_sum_data = [
        ["Toàn bộ Core Backend (Domain + Service + Repository)", "74.3% (1.627 / 2.189)", "51.0% (275 / 539)", "74.1% (7.796 / 10.527)", "Đạt yêu cầu"],
        ["8 Đơn vị mã nguồn trọng yếu (Chương 4)", "100%", "100%", "100%", "Đạt yêu cầu"],
        ["Toàn bộ Dự án (Gồm cả mã giao diện đồ họa)", "32.7% (1.627 / 4.978)", "22.9% (275 / 1.200)", "30.4% (7.796 / 25.674)", "Được kiểm thử tại tầng giao diện (Playwright E2E)"]
    ]
    t2 = doc.add_table(rows=len(cov_sum_data) + 1, cols=len(cov_sum_headers))
    t2.style = 'Table Grid'
    format_custom_table(t2, [2.2, 1.3, 1.3, 1.3, 1.3], cov_sum_headers, cov_sum_data, hdr_bg='2E75B6', hdr_text_white=True, font_size_pt=8.5)
    
    add_p(doc, text="Chi tiết mức độ bao phủ tại các gói chức năng Backend then chốt được thể hiện trong Bảng 1.3:")
    
    # Bảng 1.3: Chi tiết từng package
    p_cap3 = doc.add_paragraph()
    p_cap3.paragraph_format.space_before = Pt(8)
    p_cap3.paragraph_format.space_after = Pt(4)
    r_cap3 = p_cap3.add_run("Bảng 1.3: Chi tiết tỷ lệ bao phủ mã nguồn theo từng gói chức năng Backend")
    r_cap3.font.name = 'Times New Roman'
    r_cap3.font.size = Pt(11)
    r_cap3.font.bold = True
    r_cap3.font.italic = True
    
    pkg_headers = ["Gói chức năng (Package)", "Line Coverage", "Branch Coverage", "Instruction Cov.", "Nội dung kiểm thử chính"]
    pkg_data = [
        ["Quản lý cấu hình & kết nối", "90.9%", "75.0%", "93.0%", "Quản lý đối tượng cấu hình ứng dụng và kết nối dữ liệu"],
        ["Ngữ cảnh hệ thống (Context)", "84.8%", "57.1%", "85.0%", "Khởi tạo ngữ cảnh làm việc và nạp các kho dữ liệu"],
        ["Kho lưu trữ dữ liệu (Infrastructure)", "82.4%", "60.0%", "83.0%", "Thao tác truy xuất CSDL SQLite, Repository và nhật ký kiểm toán"],
        ["Khởi tạo danh mục sản phẩm", "77.8%", "n/a", "76.0%", "Khởi tạo đối tượng đồ uống theo từng nhóm danh mục"],
        ["Mô hình dữ liệu (Domain Model)", "77.2%", "33.3%", "64.0%", "Các thực thể nghiệp vụ Đơn hàng, Món, Nguyên liệu, Người dùng"],
        ["Dịch vụ nghiệp vụ (Service)", "75.7%", "57.8%", "72.0%", "Xử lý nghiệp vụ Đơn hàng, Thanh toán, Thực đơn, Kho nguyên liệu"],
        ["Tùy biến đồ uống (Topping & Size)", "66.7%", "n/a", "57.0%", "Tính toán lũy tiến đơn giá theo tùy biến sản phẩm"],
        ["Quản lý trạng thái đơn hàng", "63.3%", "n/a", "57.0%", "Luân chuyển và kiểm soát ràng buộc trạng thái đơn hàng"],
        ["Hệ thống thông báo sự kiện", "53.3%", "66.7%", "55.0%", "Cơ chế phát thông báo sự kiện giữa các phân hệ"],
        ["Tích hợp cổng thanh toán", "53.8%", "25.0%", "48.0%", "Xử lý giao dịch thanh toán điện tử và bắt lỗi ngoại lệ"],
        ["Chính sách chiết khấu khuyến mãi", "41.7%", "35.7%", "46.0%", "Tính toán tỷ lệ giảm giá và chiết khấu đơn hàng"]
    ]
    t3 = doc.add_table(rows=len(pkg_data) + 1, cols=len(pkg_headers))
    t3.style = 'Table Grid'
    format_custom_table(t3, [2.3, 0.9, 1.0, 1.0, 2.2], pkg_headers, pkg_data, hdr_bg='1F4E79', hdr_text_white=True, font_size_pt=8.5)

    # -------------------------------------------------------------------------
    # 1.2.4. Kết luận
    # -------------------------------------------------------------------------
    add_h1(doc, "1.2.4. Đánh giá kết quả thực thi kiểm thử Backend")
    add_p(doc, text="Kết quả thực thi kiểm thử tự động phân hệ Backend mang lại các giá trị cụ thể sau:")
    add_p(doc, text="Tỷ lệ đạt tuyệt đối: Toàn bộ 24/24 ca kiểm thử đều vượt qua (tỷ lệ thành công 100%), xác nhận tính đúng đắn của các quy tắc nghiệp vụ.", bold_prefix="1. ", indent=True)
    add_p(doc, text="Độ bao phủ đáp ứng yêu cầu: Tầng xử lý logic nghiệp vụ cốt lõi đạt 74.3% độ bao phủ dòng lệnh và 51.0% độ bao phủ nhánh rẽ, đảm bảo hầu hết các luồng xử lý chính và ngoại lệ đã được kiểm chứng.", bold_prefix="2. ", indent=True)
    add_p(doc, text="Phát hiện và ngăn chặn sớm sai sót: Các lỗi nhập liệu sai, thiếu hụt tồn kho nguyên liệu và xung đột trạng thái đơn hàng được kiểm soát và xử lý chặt chẽ tại tầng Backend trước khi tích hợp với giao diện người dùng.", bold_prefix="3. ", indent=True)

    doc.save(OUT_DOCX)
    print(f"Successfully generated updated academic Word document at: {OUT_DOCX}")

if __name__ == '__main__':
    build_docx()
