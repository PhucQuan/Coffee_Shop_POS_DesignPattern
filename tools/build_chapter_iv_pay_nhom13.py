import math
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT_DOC = ROOT / "docs" / "Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx"
ASSETS_DIR = ROOT / "build" / "chapter-iv-diagrams"
ASSETS_DIR.mkdir(parents=True, exist_ok=True)

FONT_PATH = "C:/Windows/Fonts/arial.ttf"
FONT_BOLD_PATH = "C:/Windows/Fonts/arialbd.ttf"

def get_font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD_PATH if bold else FONT_PATH, size)

def draw_grid(draw, w, h, step=25):
    for x in range(0, w, step):
        draw.line([(x, 0), (x, h)], fill="#f0f2f5", width=1)
    for y in range(0, h, step):
        draw.line([(0, y), (w, y)], fill="#f0f2f5", width=1)

def draw_arrow(draw, start, end, fill="#333333", width=2, text=None, text_offset=(10, -5)):
    x1, y1 = start
    x2, y2 = end
    draw.line([start, end], fill=fill, width=width)
    ang = math.atan2(y2 - y1, x2 - x1)
    arrow_len = 10
    p1 = (x2 - arrow_len * math.cos(ang + 0.4), y2 - arrow_len * math.sin(ang + 0.4))
    p2 = (x2 - arrow_len * math.cos(ang - 0.4), y2 - arrow_len * math.sin(ang - 0.4))
    draw.polygon([(x2, y2), p1, p2], fill=fill)
    if text:
        mx = (x1 + x2) / 2 + text_offset[0]
        my = (y1 + y2) / 2 + text_offset[1]
        draw.text((mx, my), text, fill="#c0392b" if "T" in text else "#2980b9", font=get_font(12, True))

# Common layout for CFG and Variable lifecycle graphs
# Vertical layout similar to Draw.io / Nhom13
# Width: 400, Height: 900
POS = {
    1: (200, 60),
    2: (110, 130),
    3: (200, 130),
    4: (200, 200),
    5: (200, 270),
    6: (200, 340),
    7: (200, 410),
    8: (200, 480),
    9: (120, 550),
    10: (120, 620),
    11: (280, 580),
    12: (200, 690),
    13: (200, 760),
    14: (120, 830),
    15: (200, 900),
}

def draw_base_cfg(d, highlight_nodes=None, var_defs=None, var_uses=None, show_kill_node=False, kill_from_nodes=None):
    # Connections
    # 1 -> 2
    draw_arrow(d, (POS[1][0]-15, POS[1][1]+10), (POS[2][0]+12, POS[2][1]-12))
    # 1 -> 3
    draw_arrow(d, (POS[1][0], POS[1][1]+16), (POS[3][0], POS[3][1]-16))
    # 3 -> 4
    draw_arrow(d, (POS[3][0], POS[3][1]+16), (POS[4][0], POS[4][1]-16))
    # 4 -> 5
    draw_arrow(d, (POS[4][0], POS[4][1]+16), (POS[5][0], POS[5][1]-16))
    # 4 -> 15 (bypass right)
    arc_x = 350
    d.line([(POS[4][0]+16, POS[4][1]), (arc_x, POS[4][1])], fill="#555555", width=2)
    d.line([(arc_x, POS[4][1]), (arc_x, POS[15][1])], fill="#555555", width=2)
    draw_arrow(d, (arc_x, POS[15][1]), (POS[15][0]+16, POS[15][1]))
    
    # 5 -> 6 -> 7 -> 8
    draw_arrow(d, (POS[5][0], POS[5][1]+16), (POS[6][0], POS[6][1]-16))
    draw_arrow(d, (POS[6][0], POS[6][1]+16), (POS[7][0], POS[7][1]-16))
    draw_arrow(d, (POS[7][0], POS[7][1]+16), (POS[8][0], POS[8][1]-16))
    
    # 8 -> 9
    draw_arrow(d, (POS[8][0]-12, POS[8][1]+12), (POS[9][0]+10, POS[9][1]-12))
    # 9 -> 10
    draw_arrow(d, (POS[9][0], POS[9][1]+16), (POS[10][0], POS[10][1]-16))
    # 10 -> 12
    draw_arrow(d, (POS[10][0]+10, POS[10][1]+12), (POS[12][0]-12, POS[12][1]-12))
    
    # 8 -> 11
    draw_arrow(d, (POS[8][0]+12, POS[8][1]+12), (POS[11][0]-10, POS[11][1]-12))
    # 11 -> 12
    draw_arrow(d, (POS[11][0]-10, POS[11][1]+12), (POS[12][0]+12, POS[12][1]-12))
    
    # 12 -> 13
    draw_arrow(d, (POS[12][0], POS[12][1]+16), (POS[13][0], POS[13][1]-16))
    
    # 13 -> 14
    draw_arrow(d, (POS[13][0]-12, POS[13][1]+12), (POS[14][0]+10, POS[14][1]-12))
    # 14 -> 15
    draw_arrow(d, (POS[14][0]+10, POS[14][1]+12), (POS[15][0]-12, POS[15][1]-12))
    # 13 -> 15
    draw_arrow(d, (POS[13][0], POS[13][1]+16), (POS[15][0], POS[15][1]-16))
    
    # Side green kill node (if applicable like Nhom 13)
    if show_kill_node:
        kill_pos = (50, 460)
        # arrows from exit/kill nodes to kill node
        if kill_from_nodes:
            for fn in kill_from_nodes:
                p = POS[fn]
                d.line([(p[0]-16, p[1]), (kill_pos[0]+16, kill_pos[1])], fill="#27ae60", width=1)
        # draw green kill circle
        d.ellipse((kill_pos[0]-16, kill_pos[1]-16, kill_pos[0]+16, kill_pos[1]+16), fill="#2ecc71", outline="#27ae60", width=2)
        f_k = get_font(11, True)
        d.text((kill_pos[0]-10, kill_pos[1]-7), "kill", fill="white", font=f_k)
        
    # Draw nodes 1..15
    for num in range(1, 16):
        cx, cy = POS[num]
        r = 16
        is_hl = highlight_nodes and num in highlight_nodes
        is_def = var_defs and num in var_defs
        is_use = var_uses and num in var_uses
        
        # Color matching Nhom 13 style
        fill_col = "#ffffff"
        outline_col = "#333333"
        if num == 1 and not is_use:
            fill_col = "#2ecc71" # Start node green in Nhom 13
            outline_col = "#27ae60"
        elif is_def:
            fill_col = "#ffeaa7" # Highlight def
            outline_col = "#d35400"
        elif is_use:
            fill_col = "#dff9fb" # Highlight use
            outline_col = "#0984e3"
        elif num == 2:
            fill_col = "#ffcccc" # Error
            outline_col = "#e74c3c"
            
        d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=fill_col, outline=outline_col, width=2)
        # number
        f_num = get_font(12, True)
        bbox = d.textbbox((0, 0), str(num), font=f_num)
        w = bbox[2] - bbox[0]
        h = bbox[3] - bbox[1]
        text_color = "white" if (num == 1 and not is_use) else "#2d3436"
        d.text((cx - w/2, cy - h/2 - 2), str(num), fill=text_color, font=f_num)
        
        # Annotation text on the right if specified
        ann = ""
        if is_def and is_use:
            ann = f"d, u"
        elif is_def:
            ann = f"d"
        elif is_use:
            ann = f"u"
        if ann:
            d.text((cx + r + 6, cy - 7), ann, fill="#d63031" if "d" in ann else "#0984e3", font=get_font(11, True))

def generate_all_images():
    W, H = 400, 960
    
    # 1. CFG
    im_cfg = Image.new("RGB", (W, H), "white")
    d_cfg = ImageDraw.Draw(im_cfg)
    draw_grid(d_cfg, W, H)
    draw_base_cfg(d_cfg)
    cfg_file = ASSETS_DIR / "pay_cfg_nhom13.png"
    im_cfg.save(cfg_file)
    print("Saved:", cfg_file)
    
    # 2. DFG general
    im_dfg = Image.new("RGB", (550, 960), "white")
    d_dfg = ImageDraw.Draw(im_dfg)
    draw_grid(d_dfg, 550, 960)
    # nodes with general def/use annotations
    draw_base_cfg(d_dfg, var_defs=[1, 3, 5], var_uses=[1, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 14, 15])
    dfg_file = ASSETS_DIR / "pay_dfg_nhom13.png"
    im_dfg.save(dfg_file)
    print("Saved:", dfg_file)
    
    # 3. Variable $order lifecycle graph
    im_order = Image.new("RGB", (W, H), "white")
    d_order = ImageDraw.Draw(im_order)
    draw_grid(d_order, W, H)
    draw_base_cfg(d_order, var_defs=[1], var_uses=[1, 3, 5, 6, 8, 9, 10, 11, 12, 14], show_kill_node=True, kill_from_nodes=[2, 15])
    f_order = ASSETS_DIR / "pay_order_lifecycle.png"
    im_order.save(f_order)
    print("Saved:", f_order)
    
    # 4. Variable $gateway lifecycle graph
    im_gw = Image.new("RGB", (W, H), "white")
    d_gw = ImageDraw.Draw(im_gw)
    draw_grid(d_gw, W, H)
    draw_base_cfg(d_gw, var_defs=[1], var_uses=[3, 5], show_kill_node=True, kill_from_nodes=[2, 15])
    f_gw = ASSETS_DIR / "pay_gateway_lifecycle.png"
    im_gw.save(f_gw)
    print("Saved:", f_gw)
    
    # 5. Variable $result lifecycle graph
    im_res = Image.new("RGB", (W, H), "white")
    d_res = ImageDraw.Draw(im_res)
    draw_grid(d_res, W, H)
    draw_base_cfg(d_res, var_defs=[3], var_uses=[4, 5, 15], show_kill_node=True, kill_from_nodes=[15])
    f_res = ASSETS_DIR / "pay_result_lifecycle.png"
    im_res.save(f_res)
    print("Saved:", f_res)
    
    # 6. Variable $payment lifecycle graph
    im_pay = Image.new("RGB", (W, H), "white")
    d_pay = ImageDraw.Draw(im_pay)
    draw_grid(d_pay, W, H)
    draw_base_cfg(d_pay, var_defs=[5], var_uses=[6, 7], show_kill_node=True, kill_from_nodes=[15])
    f_pay = ASSETS_DIR / "pay_payment_lifecycle.png"
    im_pay.save(f_pay)
    print("Saved:", f_pay)

def shade(cell, color="2C3E50"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)

def add_table_nhom13(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, v in enumerate(headers):
        cell = t.rows[0].cells[i]
        r = cell.paragraphs[0].add_run(v)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(cell, "2C3E50")
    for row_data in rows:
        row = t.add_row()
        for i, val in enumerate(row_data):
            c = row.cells[i]
            p = c.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
    doc.add_paragraph()
    return t

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

def add_image_centered(doc, img_path, width_cm=7.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(img_path), width=Cm(width_cm))

def build_full_docx():
    doc = Document()
    sec = doc.sections[0]
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Cm(2))
        
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11.5)

    # Tiêu đề Chương IV y hệt Nhóm 13
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Chương IV. Thực hiện kiểm thử hộp trắng 8 đơn vị mã nguồn")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    
    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run("I.  Hàm thứ nhất:")
    r_h1.bold = True
    r_h1.font.name = "Times New Roman"
    r_h1.font.size = Pt(13)
    
    code_text = (
        'public function pay($order, $gateway)\n'
        '{\n'
        '    if ($order->getPayment() !== null && $order->getPayment()->getStatus() === "SUCCESS") (1) {\n'
        '        die("Lỗi: Đơn hàng đã được thanh toán."); (2)\n'
        '    }\n\n'
        '    $result = $gateway->processPayment($order->getTotalAmount()); (3)\n\n'
        '    if ($result->isSuccess()) (4) {\n'
        '        $payment = new Payment($order->getId(), $gateway->getGatewayName(), $order->getTotalAmount(), $result->getTransactionCode(), "SUCCESS"); (5)\n\n'
        '        $order->setPayment($payment); (6)\n'
        '        self::$repository->savePayment($payment); (7)\n\n'
        '        if ($order->getStatus() === "PENDING") (8) {\n'
        '            self::$inventoryService->deductForOrder($order); (9)\n'
        '            $order->getState()->sendToKitchen($order); (10)\n'
        '        } else {\n'
        '            $order->getState()->pay($order); (11)\n'
        '        }\n\n'
        '        self::$repository->saveOrder($order); (12)\n\n'
        '        if (self::$publisher !== null) (13) {\n'
        '            self::$publisher->notifyObservers($order, $order->getStatus()); (14)\n'
        '        }\n'
        '    }\n\n'
        '    return $result; (15)\n'
        '}'
    )
    p_code = doc.add_paragraph()
    p_code.paragraph_format.left_indent = Cm(0.5)
    r_c = p_code.add_run(code_text)
    r_c.font.name = "Consolas"
    r_c.font.size = Pt(9)
    
    # 1. CFG Image
    add_image_centered(doc, ASSETS_DIR / "pay_cfg_nhom13.png", width_cm=7.2)
    
    p_deg = doc.add_paragraph()
    r_deg = p_deg.add_run("Đồ thị có 4 nút quyết định nhị phân => độ phức tạp C = 4 + 1 = 5")
    r_deg.font.name = "Times New Roman"
    
    p_paths = doc.add_paragraph()
    r_paths = p_paths.add_run("5 đường thi hành tuyến tính:\n"
                              "1->2\n"
                              "1->3->4->15\n"
                              "1->3->4->5->6->7->8->9->10->12->13->15\n"
                              "1->3->4->5->6->7->8->11->12->13->15\n"
                              "1->3->4->5->6->7->8->9->10->12->13->14->15")
    r_paths.font.name = "Times New Roman"
    
    # Table test cases
    headers_stt = ["STT", "Đường thi hành", "Bộ dữ liệu kiểm thử"]
    rows_stt = [
        ["1", "1->2", "($order: payment.status='SUCCESS', $gateway: FakeGateway)"],
        ["2", "1->3->4->15", "($order: status='PENDING', total=50000, $result: isSuccess=false)"],
        ["3", "1->3->4->5->6->7->8->9->10->12->13->15", "($order: status='PENDING', total=50000, $result: isSuccess=true, $publisher=null)"],
        ["4", "1->3->4->5->6->7->8->11->12->13->15", "($order: status='READY', total=50000, $result: isSuccess=true, $publisher=null)"],
        ["5", "1->3->4->5->6->7->8->9->10->12->13->14->15", "($order: status='PENDING', total=50000, $result: isSuccess=true, $publisher!=null)"]
    ]
    add_table_nhom13(doc, headers_stt, rows_stt)
    
    p_dfg_title = doc.add_paragraph()
    r_dfg_title = p_dfg_title.add_run("Đồ thị dòng dữ liệu:")
    r_dfg_title.bold = True
    r_dfg_title.font.name = "Times New Roman"
    
    # 2. DFG Image
    add_image_centered(doc, ASSETS_DIR / "pay_dfg_nhom13.png", width_cm=8.5)
    
    # 3. Biến $order
    p_v1 = doc.add_paragraph()
    p_v1.add_run("Kiểm thử đời sống biến $order:").bold = True
    add_image_centered(doc, ASSETS_DIR / "pay_order_lifecycle.png", width_cm=6.8)
    add_bullet(doc, "Kịch bản 1: ~duk")
    add_bullet(doc, "2: ~duuk")
    add_bullet(doc, "3: ~duuuuuuuuk")
    add_bullet(doc, "4: ~duuuuuuuk")
    add_bullet(doc, "5: ~duuuuuuuuuk")
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.left_indent = Cm(0.8)
    p_c1.add_run("Cả 5 kịch bản đều không chứa cặp đôi nào hoạt động bất thường.")
    
    # 4. Biến $gateway
    p_v2 = doc.add_paragraph()
    p_v2.add_run("Kiểm thử dòng đời biến $gateway:").bold = True
    add_image_centered(doc, ASSETS_DIR / "pay_gateway_lifecycle.png", width_cm=6.8)
    add_bullet(doc, "Kịch bản 1: ~dk")
    add_bullet(doc, "2: ~duk")
    add_bullet(doc, "3: ~duuk")
    add_bullet(doc, "4: ~duuk (giống 3)")
    add_bullet(doc, "5: ~duuk (giống 3)")
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.left_indent = Cm(0.8)
    p_c2.add_run("Cả 5 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")
    
    # 5. Biến $result
    p_v3 = doc.add_paragraph()
    p_v3.add_run("Kiểm thử dòng đời biến $result:").bold = True
    add_image_centered(doc, ASSETS_DIR / "pay_result_lifecycle.png", width_cm=6.8)
    add_bullet(doc, "Kịch bản 1: không đi qua câu lệnh (3)")
    add_bullet(doc, "2: ~duuk")
    add_bullet(doc, "3: ~duuuk")
    add_bullet(doc, "4: ~duuuk (giống 3)")
    add_bullet(doc, "5: ~duuuk (giống 3)")
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.left_indent = Cm(0.8)
    p_c3.add_run("Cả 5 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")
    
    # 6. Biến $payment
    p_v4 = doc.add_paragraph()
    p_v4.add_run("Kiểm thử dòng đời biến $payment:").bold = True
    add_image_centered(doc, ASSETS_DIR / "pay_payment_lifecycle.png", width_cm=6.8)
    add_bullet(doc, "Kịch bản 1: không đi qua câu lệnh (5)")
    add_bullet(doc, "2: không đi qua câu lệnh (5)")
    add_bullet(doc, "3: ~duuk")
    add_bullet(doc, "4: ~duuk (giống 3)")
    add_bullet(doc, "5: ~duuk (giống 3)")
    p_c4 = doc.add_paragraph()
    p_c4.paragraph_format.left_indent = Cm(0.8)
    p_c4.add_run("Cả 5 kịch bản trên đều không chứa cặp đôi nào hoạt động bất thường.")

    # Try saving to primary file, and always save to Nhom13 specific copy
    target_v2 = ROOT / "docs" / "Chuong_IV_KiemThuHopTrang_CoffeeShopPOS_Nhom13.docx"
    doc.save(target_v2)
    print("Saved to:", target_v2)
    try:
        doc.save(OUT_DOC)
        print("Full docx updated successfully at:", OUT_DOC)
    except PermissionError:
        print("Note: Primary file is currently open in Word. Saved to Chuong_IV_KiemThuHopTrang_CoffeeShopPOS_Nhom13.docx")

if __name__ == "__main__":
    generate_all_images()
    build_full_docx()
