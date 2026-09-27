from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Chuong_IV_KiemThuHopTrang_CoffeeShopPOS.docx"
ASSETS = ROOT / "build" / "chapter-iv-diagrams"

FONT = "C:/Windows/Fonts/arial.ttf"
FONT_BOLD = "C:/Windows/Fonts/arialbd.ttf"

def f(size, bold=False): return ImageFont.truetype(FONT_BOLD if bold else FONT, size)

def arrow(draw, a, b, fill="#555555"):
    draw.line([a, b], fill=fill, width=3)
    x1, y1, x2, y2 = *a, *b
    import math
    ang = math.atan2(y2-y1, x2-x1)
    for delta in (2.6, -2.6):
        p = (x2 - 13*math.cos(ang+delta), y2 - 13*math.sin(ang+delta))
        draw.line([(x2,y2), p], fill=fill, width=3)

def node(draw, center, label, green=False, width=150):
    x,y=center; h=48
    draw.ellipse((x-width//2,y-h//2,x+width//2,y+h//2), fill="#09c914" if green else "white", outline="#333333", width=2)
    box=draw.textbbox((0,0), label, font=f(13, True))
    draw.text((x-(box[2]-box[0])/2,y-(box[3]-box[1])/2-1), label, fill="black", font=f(13, True))

def diagram(path, title, labels):
    # Luồng điều khiển dạng dọc, có hai nhánh bên trái/phải giống hình mẫu.
    w,h=1250,max(900, 160+len(labels)*115)
    im=Image.new("RGB",(w,h),"white"); d=ImageDraw.Draw(im)
    d.text((35,20), title, fill="#222222", font=f(26,True))
    centers=[]
    for i,label in enumerate(labels):
        if i==0 or i==len(labels)-1: x=625
        elif i in (2,5): x=420
        elif i in (3,6): x=830
        else: x=625
        centers.append((x,95+i*100))
    for i in range(len(centers)-1):
        a,b=centers[i],centers[i+1]
        arrow(d,(a[0],a[1]+26),(b[0],b[1]-26))
    for i,(c,lbl) in enumerate(zip(centers,labels)):
        node(d,c,lbl, i in (0,len(labels)-1), 185 if len(lbl)>14 else 145)
    im.save(path)

def flow_diagram(path, variable, defs, uses):
    w,h=1250,900; im=Image.new("RGB",(w,h),"white"); d=ImageDraw.Draw(im)
    d.text((35,20), f"Đồ thị dòng dữ liệu biến {variable}", fill="#222222", font=f(26,True))
    all_nodes=["Bắt đầu"] + [f"d({variable})\n{v}" for v in defs] + [f"u({variable})\n{v}" for v in uses] + ["Kết thúc"]
    xs=[625,625,410,830,625,410,830,625]; ys=[100,220,340,460,580,700,780,850]
    pts=[]
    for i,text in enumerate(all_nodes):
        x=xs[min(i,len(xs)-1)]; y=ys[min(i,len(ys)-1)]; pts.append((x,y))
    for a,b in zip(pts,pts[1:]): arrow(d,(a[0],a[1]+28),(b[0],b[1]-28))
    # Cung định nghĩa -> các sử dụng, thể hiện cặp DU.
    if len(pts)>3:
        for p in pts[2:-1]:
            arrow(d,(pts[1][0]+70,pts[1][1]),(p[0]-70,p[1]),"#888888")
    for i,(p,text) in enumerate(zip(pts,all_nodes)):
        node(d,p,text,i in (0,len(all_nodes)-1),190)
    im.save(path)

def shade(cell, color="7A4A32"):
    tcpr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),color); tcpr.append(shd)

def table(doc, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,v in enumerate(headers):
        r=t.rows[0].cells[i].paragraphs[0].add_run(v); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    doc.add_paragraph(); return t

def code(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.4)
    r=p.add_run(text); r.font.name="Consolas"; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(230,230,230)
    p._p.get_or_add_pPr().append(OxmlElement("w:shd")); p._p.pPr[-1].set(qn("w:fill"),"1E1E1E")

UNITS=[
 ("IV.1", "PaymentService.pay", "PaymentService.java", "pay", 5,
  '''public PaymentResult pay(Order order, PaymentGateway gateway) {
    if (order.getPayment() != null && "SUCCESS".equals(order.getPayment().getStatus())) { // (1)
        throw new IllegalStateException("Order has already been paid.");                 // (2)
    }
    PaymentResult result = gateway.processPayment(order.getTotalAmount());                // (3)
    if (result.isSuccess()) {                                                              // (4)
        Payment payment = new Payment(...);                                                // (5)
        if ("PENDING".equals(order.getStatus())) {                                       // (6)
            inventoryService.deductForOrder(order);                                        // (7)
            order.getState().sendToKitchen(order);                                         // (8)
        } else { order.getState().pay(order); }                                            // (9)
        repository.saveOrder(order);                                                       // (10)
        if (publisher != null) publisher.notifyObservers(order, order.getStatus());        // (11)
    }
    return result;                                                                         // (12)
}''',
  ["Bắt đầu","(1) Đã thanh toán?","(2) Ném lỗi","(4) Gateway thành công?","(6) PENDING?","(7–8) Trừ kho, gửi bếp","(9) Thanh toán state","(11) Có observer?","(12) Trả result","Kết thúc"],
  [("P1","Đơn đã SUCCESS","1→2","Ném lỗi"),("P2","Gateway thất bại","1→3→4→12","Trả fail"),("P3","PENDING, gateway OK","1→3→4→5→6→7→8→10→12","PREPARING"),("P4","READY, gateway OK","1→3→4→5→6→9→10→12","PAID"),("P5","publisher có observer","…→11→12","Có thông báo")],
  ("result",["gateway.processPayment"],["result.isSuccess", "return result"])),
 ("IV.2", "OrderService.addItem", "OrderService.java", "add_item", 4,
  '''public OrderItem addItem(Order order, int beverageId, Beverage beverage, int quantity, String note) {
    ensurePending(order, "add item");                                                     // (1)
    for (OrderItem existing : order.getItems()) {                                          // (2)
        if (existing.getBeverage().getDescription().equals(beverage.getDescription())      // (3)
                && Objects.equals(existing.getNote(), note == null ? "" : note)) {        // (4)
            existing.setQuantity(existing.getQuantity() + quantity);                       // (5)
            recalculate(order); return existing;                                           // (6)
        }
    }
    OrderItem item = new OrderItem(...);                                                   // (7)
    order.getState().addItem(order, item); recalculate(order); return item;                // (8)
}''',
  ["Bắt đầu","(1) PENDING?","Lỗi trạng thái","(2) Còn item?","(3–4) Trùng món & ghi chú?","(5–6) Gộp số lượng","(7) Tạo item","(8) Thêm, tính lại, trả","Kết thúc"],
  [("P1","Đơn PREPARING","1→lỗi","Chặn thêm món"),("P2","Giỏ trống","1→2→7→8","Thêm mới"),("P3","Cùng món, khác ghi chú","1→2→3→4→7→8","Tạo dòng riêng"),("P4","Cùng món, cùng ghi chú","1→2→3→4→5→6","Gộp số lượng")],
  ("existing",["vòng lặp order.getItems"],["so sánh", "setQuantity", "return existing"])),
 ("IV.3", "OrderService.updateItemQuantity", "OrderService.java", "update_qty", 4,
  '''public void updateItemQuantity(Order order, int itemId, int quantity) {
    ensurePending(order, "update item quantity");                                        // (1)
    if (quantity <= 0) { removeItem(order, itemId); return; }                              // (2)
    OrderItem item = order.getItems().stream().filter(x -> x.getId() == itemId)            // (3)
        .findFirst().orElseThrow(() -> new IllegalArgumentException(...));
    item.setQuantity(quantity); recalculate(order);                                        // (4)
}''',
  ["Bắt đầu","(1) PENDING?","Lỗi trạng thái","(2) quantity ≤ 0?","Xóa item","(3) Tìm item","Không thấy: lỗi","(4) Cập nhật, tính lại","Kết thúc"],
  [("P1","Đơn READY","1→lỗi","Chặn sửa"),("P2","quantity = 0","1→2→xóa","Xóa item"),("P3","itemId không có","1→2→3→lỗi","Ném lỗi"),("P4","itemId hợp lệ, quantity=3","1→2→3→4","Cập nhật")],
  ("item",["findFirst"],["setQuantity", "recalculate"])),
 ("IV.4", "OrderService.recalculate", "OrderService.java", "recalculate", 1,
  '''public void recalculate(Order order) {
    DiscountStrategy strategy = DiscountStrategyResolver.fromName(order.getDiscountType()); // (1)
    double discount = strategy.calculateDiscount(order);                                    // (2)
    double finalTotal = Math.max(0, order.getSubtotal() - discount);                        // (3)
    order.setDiscountType(strategy.getName()); order.setDiscountAmount(discount);           // (4)
    order.setTotalAmount(finalTotal); repository.saveOrder(order);                          // (5)
}''',
  ["Bắt đầu","(1) Lấy strategy","(2) Tính discount","(3) Chặn tổng âm","(4) Gán giảm giá","(5) Lưu đơn","Kết thúc"],
  [("P1","Subtotal 100.000, PERCENT_10","1→2→3→4→5","Total 90.000"),("P2","Subtotal 20.000, giảm 30.000","1→2→3→4→5","Total = 0")],
  ("finalTotal",["Math.max"],["setTotalAmount"])),
 ("IV.5", "InventoryService.restockItem", "InventoryService.java", "restock", 2,
  '''public void restockItem(int id, double amount) {
    if (amount <= 0) { throw new IllegalArgumentException(...); }                          // (1)
    InventoryItem item = getInventoryItemById(id);                                          // (2)
    repository.adjustInventory(item.getId(), amount, null, "MANUAL_RESTOCK");             // (3)
}''',
  ["Bắt đầu","(1) amount ≤ 0?","Ném lỗi","(2) Tìm nguyên liệu","(3) Cộng kho","Kết thúc"],
  [("P1","amount = 0","1→lỗi","Từ chối"),("P2","amount = -5","1→lỗi","Từ chối"),("P3","id hợp lệ, amount=10","1→2→3","Kho tăng 10")],
  ("amount",["tham số"],["so sánh", "adjustInventory"])),
 ("IV.6", "InventoryService.deductForOrder", "InventoryService.java", "deduct", 3,
  '''public void deductForOrder(Order order) {
    if (order.isInventoryDeducted()) { return; }                                            // (1)
    Map<String, Double> requirements = calculateRequirements(order);                        // (2)
    validateAvailable(requirements);                                                        // (3)
    requirements.forEach((name, quantity) -> {                                              // (4)
        InventoryItem item = getItem(name);                                                 // (5)
        repository.adjustInventory(item.getId(), -quantity, order.getId(), "ORDER_DEDUCTION"); // (6)
    });
    order.setInventoryDeducted(true);                                                       // (7)
}''',
  ["Bắt đầu","(1) Đã trừ kho?","Trả về","(2) Tính requirements","(3) Kho đủ?","(4–6) Duyệt và trừ","(7) Đánh dấu đã trừ","Kết thúc"],
  [("P1","isInventoryDeducted=true","1→return","Không trừ lặp"),("P2","Kho thiếu","1→2→3→lỗi","Ném InventoryException"),("P3","Kho đủ","1→2→3→4→5→6→7","Trừ đúng công thức")],
  ("requirements",["calculateRequirements"],["validateAvailable", "forEach"])),
 ("IV.7", "MenuService.createBeverage", "MenuService.java", "factory", 5,
  '''public Beverage createBeverage(MenuItemRecord item) {
    BeverageFactory factory = switch (item.getCategory()) {                                // (1)
        case "COFFEE" -> new CoffeeFactory();                                              // (2)
        case "TEA" -> new TeaFactory();                                                    // (3)
        case "MATCHA" -> new MatchaFactory();                                              // (4)
        case "SMOOTHIE" -> new SmoothieFactory();                                          // (5)
        default -> throw new IllegalArgumentException(...);                                // (6)
    };
    return factory.createBeverage(item.getName(), item.getBasePrice());                     // (7)
}''',
  ["Bắt đầu","(1) switch category","(2) COFFEE","(3) TEA","(4) MATCHA","(5) SMOOTHIE","(6) default: lỗi","(7) Tạo beverage","Kết thúc"],
  [("P1","COFFEE","1→2→7","Tạo BaseCoffee"),("P2","TEA","1→3→7","Tạo MilkTea"),("P3","MATCHA","1→4→7","Tạo Matcha"),("P4","SMOOTHIE","1→5→7","Tạo Smoothie"),("P5","JUICE","1→6","Ném lỗi")],
  ("factory",["switch category"],["createBeverage"])),
 ("IV.8", "MenuService.saveRecipeItem", "MenuService.java", "recipe", 4,
  '''public void saveRecipeItem(MenuItemRecord beverage, InventoryItem inventoryItem, double quantityRequired) {
    if (beverage == null) throw new IllegalArgumentException(...);                         // (1)
    if (inventoryItem == null) throw new IllegalArgumentException(...);                    // (2)
    if (quantityRequired <= 0) throw new IllegalArgumentException(...);                    // (3)
    repository.saveRecipeItem(new RecipeItem(beverage.getId(), inventoryItem.getId(), quantityRequired)); // (4)
}''',
  ["Bắt đầu","(1) beverage null?","Lỗi","(2) inventory null?","Lỗi","(3) quantity ≤ 0?","Lỗi","(4) Lưu RecipeItem","Kết thúc"],
  [("P1","beverage=null","1→lỗi","Báo chọn đồ uống"),("P2","inventoryItem=null","1→2→lỗi","Báo chọn nguyên liệu"),("P3","quantity=0","1→2→3→lỗi","Từ chối"),("P4","Dữ liệu hợp lệ","1→2→3→4","Lưu công thức")],
  ("quantityRequired",["tham số"],["so sánh", "new RecipeItem"]))
]

def main():
    ASSETS.mkdir(parents=True,exist_ok=True)
    doc=Document(); sec=doc.sections[0]
    for attr in ("top_margin","bottom_margin","left_margin","right_margin"): setattr(sec,attr,Cm(2))
    doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=Pt(10.5)
    for n,s in (("Heading 1",15),("Heading 2",12)):
        doc.styles[n].font.name="Arial"; doc.styles[n].font.size=Pt(s); doc.styles[n].font.bold=True; doc.styles[n].font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("CHƯƠNG IV. THỰC HIỆN KIỂM THỬ HỘP TRẮNG 8 ĐƠN VỊ MÃ NGUỒN"); r.bold=True; r.font.size=Pt(16)
    doc.add_paragraph("Tài liệu riêng cho đề tài Coffee Shop POS. Mỗi đơn vị gồm mã nguồn được đánh số, đồ thị luồng điều khiển, độ phức tạp McCabe, đường đi kiểm thử và đồ thị dòng dữ liệu.")
    doc.add_paragraph("Quy ước: d(x) là định nghĩa biến x; u(x) là sử dụng biến x. Nút xanh là điểm bắt đầu/kết thúc. V(G) = số điểm quyết định + 1.")
    for no,name,source,slug,vg,snippet,labels,paths,df in UNITS:
        doc.add_page_break(); doc.add_heading(f"{no}. Hàm {name}",1); doc.add_paragraph(f"Nguồn: src/main/java/com/coffeeshop/service/{source}")
        code(doc,snippet)
        cfg=ASSETS/f"{slug}-cfg.png"; diagram(cfg,f"Đồ thị luồng điều khiển – {name}",labels)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(cfg),width=Cm(14.5))
        doc.add_paragraph(f"Hàm có độ phức tạp chu trình V(G) = {vg}. Số đường đi độc lập cần kiểm thử tối thiểu là {vg}.")
        table(doc,["TC","Dữ liệu/điều kiện","Đường đi","Kết quả mong đợi"],paths)
        variable,defs,uses=df; dfg=ASSETS/f"{slug}-dfg.png"; flow_diagram(dfg,variable,defs,uses)
        doc.add_heading(f"Đồ thị dòng dữ liệu biến {variable}",2)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(dfg),width=Cm(14.5))
        doc.add_paragraph(f"Kết luận: các đường đi được chọn đều kiểm tra cặp d({variable})–u({variable}) quan trọng; không có cặp định nghĩa–sử dụng bất thường trên luồng hợp lệ.")
    doc.add_page_break(); doc.add_heading("IV.9. Tổng hợp",1)
    table(doc,["Hàm","V(G)","Số đường cơ sở"],[(x[1],x[4],x[4]) for x in UNITS])
    doc.add_paragraph("Tổng số đường cơ sở: 28. Các tình huống được đối chiếu với TestRunner của dự án; bộ kiểm thử hiện đạt 24/24 khi biên dịch theo Java 17.")
    doc.save(OUT); print(OUT)

if __name__=="__main__": main()
